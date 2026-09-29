"""Use the verified internship worksheet as a short-lived job feed."""

from __future__ import annotations

import os
from datetime import date
from pathlib import Path
from typing import Iterator
from urllib.parse import urlparse

from openpyxl import load_workbook

from bot.search.base import BaseSearcher, RawJob


class WorkbookSearcher(BaseSearcher):
    def search(self, criteria, page=None, include_stale: bool = False) -> Iterator[RawJob]:
        """Open, recently verified, not-yet-due rows with an https URL.

        ``include_stale`` drops the status/freshness/deadline filters — used to
        learn which companies and URLs the workbook covers, never to apply.
        """
        path = Path(os.environ.get("AUTOAPPLY_WORKBOOK", ""))
        if not path.is_file():
            return
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook["Verified Opportunities"]
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows)
            for values in rows:
                record = dict(zip(headers, values))
                if not record.get("Record ID"):
                    continue
                if not include_stale:
                    if record.get("Application Status") not in {"Open", "Deadline Approaching"}:
                        continue
                    verified = _date_prefix(record.get("Date Verified"))
                    if verified is None or (date.today() - verified).days > 7:
                        continue
                    deadline = _date_prefix(record.get("Application Deadline"))
                    if deadline is not None and deadline < date.today():
                        continue
                url = str(record.get("Official Posting / Application URL") or "")
                if urlparse(url).scheme != "https" or not urlparse(url).netloc:
                    continue
                parts = [record.get(name) for name in (
                    "Recruiting Cycle", "Role / Function", "Field / Major",
                    "Class Year Eligibility", "Major Eligibility", "Other Eligibility / Notes",
                )]
                location = ", ".join(str(record.get(name)) for name in ("City", "State")
                                     if record.get(name) not in (None, "Not stated"))
                yield RawJob(
                    title=str(record.get("Program / Position") or ""),
                    company=str(record.get("Employer / Organization") or ""),
                    location=location,
                    salary=str(record.get("Pay / Stipend") or "") or None,
                    description=". ".join(str(value) for value in parts if value),
                    apply_url=url,
                    platform="workbook",
                    external_id=str(record["Record ID"]),
                    posted_at=None,
                )
        finally:
            workbook.close()


def _date_prefix(value) -> date | None:
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value)[:10])
    except (TypeError, ValueError):
        return None
