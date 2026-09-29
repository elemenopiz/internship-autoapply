"""Workbook source: the user's "verified opportunities" spreadsheet (docs/SPEC.md section 5.3).

The real sheet is hand maintained, so nothing about its layout is assumed beyond "a header row somewhere in
the first 15 rows with at least three recognisable columns". Everything is deliberately forgiving: aliases,
fuzzy headers, hyperlinks (cell links and ``=HYPERLINK()`` formulas), Excel serial dates and text dates,
free-form terms ("Summer '27", "2027 Summer"), free-form statuses, merged cells, ragged rows. A row that cannot
be understood is logged and skipped; it never aborts the ingest.

Filters (each rejected row records one primary ``RejectReason``): open, deadline not past, term ==
``search.target_term``, internship-ish, recent (``last_verified`` / ``posted`` within ``search.recent_days``).

openpyxl's streaming reader (``read_only=True``) exposes neither hyperlinks nor merged ranges, so the sheet is
read with the normal reader (``data_only=True``: cached formula results); ``=HYPERLINK()`` formulas are picked
up by a second streaming pass. Files above ``LARGE_FILE_BYTES`` fall back to streaming (no links / merges).
"""

from __future__ import annotations

import calendar
import difflib
import io
import logging
import math
import os
import re
import warnings
from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from datetime import time as dtime
from enum import StrEnum
from functools import lru_cache
from pathlib import Path
from typing import Any, Literal, NamedTuple
from urllib.parse import parse_qsl, urlsplit

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from pydantic import BaseModel, Field

from autoapply.clock import local_day
from autoapply.config import AppConfig, AppPaths
from autoapply.contracts import SourceContext
from autoapply.models import ATS, Opportunity, OpportunitySource
from autoapply.normalize import canonical_url, host_of, norm_text

_LOG = logging.getLogger("autoapply.sources.workbook")

HEADER_SCAN_ROWS = 15  # SPEC: the header row is within the first 15 rows ...
MIN_HEADER_FIELDS = 3  # ... and has at least three recognised columns
MAX_ROWS = 10_000  # a hand-made sheet never gets near this; guards against bogus dimensions
MAX_COLS = 80
LARGE_FILE_BYTES = 40 * 1024 * 1024

_TITLE_LIMIT = 300
_COMPANY_LIMIT = 200
_TEXT_LIMIT = 1_000  # location, ATS hint, ... and every ``extra`` value
_DESCRIPTION_LIMIT = 20_000
_URL_LIMIT = 4_096
_SCAN_LIMIT = 4_000  # characters of free text inspected by the term / internship heuristics


def _words(text: str, sep: str = " ") -> tuple[str, ...]:
    return tuple(w.strip() for w in text.split(sep) if w.strip())


class WorkbookError(ValueError):
    """The workbook cannot be used at all (missing, unreadable, no usable sheet or header)."""


class RejectReason(StrEnum):
    """Why a row was not turned into an opportunity (one primary reason per rejected row)."""

    MISSING_FIELDS = "missing_fields"  # has cells, but no company or no title
    CLOSED = (
        "closed"  # status says closed / filled / expired / inactive / "no" (or a closed section)
    )
    ALREADY_APPLIED = "already_applied"  # status says the user already applied
    DEADLINE_PASSED = "deadline_passed"
    WRONG_TERM = "wrong_term"  # explicit different term (Fall 2026, Summer 2028, ...)
    NOT_INTERNSHIP = "not_internship"  # senior / full-time / new-grad role without intern wording
    STALE = "stale"  # last_verified / posted older than search.recent_days
    NO_URL = "no_url"  # nothing to apply to
    ERROR = "error"  # unexpected problem while parsing the row


# ------------------------------------------------------------------------------------------------ ATS

_ATS_SUFFIXES: tuple[tuple[ATS, tuple[str, ...]], ...] = (
    (ATS.WORKDAY, ("myworkdayjobs.com", "myworkdaysite.com")),
    (ATS.GREENHOUSE, ("greenhouse.io", "grnh.se")),
    (ATS.LEVER, ("lever.co",)),
    (ATS.ASHBY, ("ashbyhq.com",)),
    (ATS.ICIMS, ("icims.com",)),
    (ATS.SMARTRECRUITERS, ("smartrecruiters.com",)),
    (ATS.TALEO, ("taleo.net",)),
    (ATS.SUCCESSFACTORS, ("successfactors.com", "successfactors.eu", "sapsf.com", "sapsf.eu")),
    (ATS.ORACLE, ("oraclecloud.com",)),
)
# Job boards / aggregators: never a direct ATS and never an employer-hosted portal.
_AGGREGATOR_SUFFIXES = _words(
    "linkedin.com lnkd.in indeed.com glassdoor.com ziprecruiter.com simplyhired.com monster.com "
    "careerbuilder.com joinhandshake.com wayup.com ripplematch.com simplify.jobs jobright.ai "
    "wellfound.com angel.co builtin.com google.com bing.com"
)
# Third-party portals we have no adapter (or enum value) for: not "employer hosted" either.
_OTHER_ATS_SUFFIXES = _words(
    "jobvite.com workable.com breezy.hr bamboohr.com applytojob.com jazz.co paylocity.com ultipro.com "
    "ukg.com adp.com paycomonline.net rippling.com dayforcehcm.com brassring.com kenexa.com pinpointhq.com "
    "recruitee.com teamtailor.com personio.de personio.com avature.net phenompeople.com eightfold.ai "
    "jobs2web.com hirevue.com cornerstoneondemand.com"
)
_CAREER_LABELS = frozenset(
    _words(
        "talent recruiting recruitment recruit hiring hire apply join joinus employment opportunities "
        "internships internship campus students student earlycareers early-careers university"
    )
)
_CAREER_PATH_SEGMENTS = frozenset(
    _words(
        "careers career jobs job opportunities positions openings internships students campus "
        "early-careers earlycareers join apply"
    )
)
_IPV4 = re.compile(r"^\d{1,3}(?:\.\d{1,3}){3}$")


def _host_in(host: str, suffixes: Iterable[str]) -> bool:
    return any(host == s or host.endswith("." + s) for s in suffixes)


def is_aggregator_url(url: str | None) -> bool:
    """True for job-board / aggregator hosts (LinkedIn, Indeed, Glassdoor, Handshake, ...)."""
    host = host_of(url)
    return bool(host) and _host_in(host, _AGGREGATOR_SUFFIXES)


def _split_url(url: str) -> tuple[list[str], dict[str, str]]:
    raw = url.strip()
    if "://" not in raw:
        raw = "https://" + raw
    try:
        parts = urlsplit(raw)
        query = {k.lower(): v for k, v in parse_qsl(parts.query, keep_blank_values=True)}
    except ValueError:
        return [], {}
    return [s.lower() for s in parts.path.split("/") if s], query


def _employer_hosted(host: str, segments: Sequence[str]) -> bool:
    if not _IPV4.match(host):
        labels = host.split(".")[:-1] or host.split(".")
        if any(lb in _CAREER_LABELS or "career" in lb or "jobs" in lb for lb in labels):
            return True
    return any(seg in _CAREER_PATH_SEGMENTS for seg in segments)


def detect_ats(url: str | None) -> ATS:
    """Classify the application system behind ``url`` from the URL alone (no I/O).

    Known ATS hosts win (Workday ``*.myworkdayjobs.com`` / ``*.myworkdaysite.com``, Greenhouse, Lever, Ashby,
    iCIMS, SmartRecruiters, Taleo, SuccessFactors, Oracle). Employer pages embedding Greenhouse
    (``?gh_jid=``) count as Greenhouse. A clearly employer-hosted careers URL (``careers.acme.com``,
    ``acme.com/careers/...``) is ``CUSTOM``; aggregators, other third-party portals and anything else are
    ``UNKNOWN``. Goes through ``normalize.host_of`` so ``*.localhost`` mock hosts behave like production ones.
    """
    if not url or not url.strip():
        return ATS.UNKNOWN
    host = host_of(url)
    if not host:
        return ATS.UNKNOWN
    for ats, suffixes in _ATS_SUFFIXES:
        if _host_in(host, suffixes):
            return ats
    segments, query = _split_url(url)
    if "gh_jid" in query:
        return ATS.GREENHOUSE
    if _host_in(host, _AGGREGATOR_SUFFIXES) or _host_in(host, _OTHER_ATS_SUFFIXES):
        return ATS.UNKNOWN
    return ATS.CUSTOM if _employer_hosted(host, segments) else ATS.UNKNOWN


_ATS_HINTS: tuple[tuple[re.Pattern[str], ATS], ...] = tuple(
    (re.compile(pattern), ats)
    for pattern, ats in (
        (r"\bworkday\b|\bmyworkday", ATS.WORKDAY),
        (r"\bgreenhouse\b", ATS.GREENHOUSE),
        (r"\blever\b", ATS.LEVER),
        (r"\bashby(?:hq)?\b", ATS.ASHBY),
        (r"\bicims\b", ATS.ICIMS),
        (r"\bsmart ?recruiters?\b", ATS.SMARTRECRUITERS),
        (r"\btaleo\b", ATS.TALEO),
        (r"\bsuccess ?factors?\b|\bsap sf\b", ATS.SUCCESSFACTORS),
        (r"\boracle\b|\borc\b", ATS.ORACLE),
        (
            r"\bcustom\b|\bemployer\b|\bcompany (?:site|website|portal)\b|\bcareers? (?:site|page|portal)\b"
            r"|\bdirect\b|\bin house\b|\bown portal\b|\bproprietary\b|\bother\b",
            ATS.CUSTOM,
        ),
    )
)


def parse_ats_hint(text: str | None) -> ATS | None:
    """Map an "ATS / Platform" cell ("Workday", "SAP SuccessFactors", "Company site") to ``ATS``; else None."""
    n = norm_text(text)
    if not n:
        return None
    for pattern, ats in _ATS_HINTS:
        if pattern.search(n):
            return ats
    return None


# ------------------------------------------------------------------------------------------------ text helpers

_PLACEHOLDERS = frozenset(
    _words("n/a na n.a. none null nil nan - -- --- — – tbd tba ? ?? unknown")
    + ("not applicable", "not available")
)
_EXCEL_ERROR = re.compile(r"^#(?:N/A|REF!|VALUE!|NAME\?|DIV/0!|NULL!|NUM!|SPILL!|CALC!)$", re.I)
_INVISIBLE = dict.fromkeys(map(ord, "​‌‍⁠﻿­"), None)
_HSPACE = re.compile(r"[ \t\f\v    　]+")


def clean_text(value: object, *, multiline: bool = False, limit: int | None = None) -> str:
    """Cell value -> tidy text: no zero-width/NBSP noise, collapsed whitespace, "" for placeholders.

    "N/A", "TBD", "-" and Excel error values (``#N/A``) become "". Booleans become "" (they carry meaning
    only for status columns). Dates render as ISO. ``multiline`` keeps line breaks (descriptions).
    """
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, datetime):
        text = value.date().isoformat()
    elif isinstance(value, date):
        text = value.isoformat()
    elif isinstance(value, float):
        text = str(int(value)) if value.is_integer() else repr(value)
    else:
        text = str(value)
    text = text.translate(_INVISIBLE)
    if multiline:
        lines = [_HSPACE.sub(" ", ln).strip() for ln in re.split(r"\r\n|\r|\n", text)]
        text = re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()
    else:
        text = _HSPACE.sub(" ", re.sub(r"[\r\n]+", " ", text)).strip()
    if not text or text.lower() in _PLACEHOLDERS or _EXCEL_ERROR.match(text):
        return ""
    if limit is not None and len(text) > limit:
        text = text[: limit - 1].rstrip() + "…"
    return text


_URL_IN_TEXT = re.compile(r"(?i)(?:https?://|www\.)[^\s<>\"'“”]+")
_BARE_URL = re.compile(
    r"(?i)^[a-z0-9][a-z0-9-]*(?:\.[a-z0-9-]+)*\.[a-z]{2,}(?::\d+)?(?:[/?#]\S*)?$"
)


def _strip_url_tail(url: str) -> str:
    url = url.rstrip(".,;:!?'\"”’")
    while url.endswith(")") and url.count("(") < url.count(")"):
        url = url[:-1].rstrip(".,;:!?")
    return url


def _valid_http_url(url: str) -> bool:
    if not url or len(url) > _URL_LIMIT or re.search(r"\s", url):
        return False
    try:
        parts = urlsplit(url)
        host = (parts.hostname or "").lower()
    except ValueError:
        return False
    if parts.scheme not in ("http", "https") or not host:
        return False
    return "." in host or host == "localhost"


def extract_url(text: object) -> str | None:
    """First usable ``http(s)`` URL in ``text``; also accepts ``www.x.com/..`` and bare ``host.tld/path``.

    Trailing punctuation is trimmed; ``mailto:`` / ``javascript:`` / relative / intra-workbook links never count.
    """
    raw = clean_text(text)
    if not raw:
        return None
    raw = raw.strip("<>#")
    if _valid_http_url(raw):
        return raw
    for match in _URL_IN_TEXT.finditer(raw):
        candidate = _strip_url_tail(match.group(0))
        if candidate.lower().startswith("www."):
            candidate = "https://" + candidate
        if _valid_http_url(candidate):
            return candidate
    if _BARE_URL.match(raw) and _valid_http_url("https://" + raw):
        return "https://" + raw
    return None


# ------------------------------------------------------------------------------------------------ dates

_MONTHS: dict[str, int] = {
    name: number
    for number, names in enumerate(
        (
            "jan january",
            "feb february",
            "mar march",
            "apr april",
            "may",
            "jun june",
            "jul july",
            "aug august",
            "sep sept september",
            "oct october",
            "nov november",
            "dec december",
        ),
        start=1,
    )
    for name in names.split()
}
_WEEKDAYS = frozenset(
    _words(
        "mon monday tue tues tuesday wed weds wednesday thu thur thurs thursday fri friday sat saturday "
        "sun sunday"
    )
)
_DATE_NOISE = frozenset(
    _words(
        "st nd rd th of the on at by as last verified posted updated added checked utc gmt est edt cst cdt "
        "mst mdt pst pdt am pm z t"
    )
)
_LEAD_NOISE = re.compile(
    r"^(?:(?:last\s+)?(?:verified|posted|updated|added|checked|confirmed)|as of|on)\s*[:\-]?\s*"
)
_TIME_OF_DAY = re.compile(r"\b\d{1,2}:\d{2}(?::\d{2}(?:\.\d+)?)?\s*(?:am|pm)?\b", re.I)
_RELATIVE = re.compile(
    r"(?:about |over |posted |verified )?(\d+)\+?\s*(hour|hr|day|d|week|wk|month|mo)s?\s+ago"
)
_RELATIVE_DAYS = {"hour": 0, "hr": 0, "day": 1, "d": 1, "week": 7, "wk": 7, "month": 30, "mo": 30}
_EXCEL_1900 = datetime(1899, 12, 30)
_YEAR_MIN, _YEAR_MAX = 1990, 2100


def _safe_date(year: int, month: int, day: int) -> date | None:
    if not (_YEAR_MIN <= year <= _YEAR_MAX):
        return None
    try:
        return date(year, month, day)
    except ValueError:
        return None


def _month_end(year: int, month: int) -> date | None:
    if not (_YEAR_MIN <= year <= _YEAR_MAX and 1 <= month <= 12):
        return None
    return date(year, month, calendar.monthrange(year, month)[1])


def _two_digit_year(value: int) -> int:
    return 2000 + value if value < 70 else 1900 + value


def _from_number(number: float, epoch: datetime) -> date | None:
    """Excel serial (1900 or 1904 epoch) or ``YYYYMMDD`` integer -> date; implausible numbers -> None."""
    if not math.isfinite(number):
        return None
    if float(number).is_integer() and 19_000_101 <= number <= 21_001_231:
        n = int(number)
        return _safe_date(n // 10_000, (n // 100) % 100, n % 100)
    if number < 1 or number > 2_958_465:
        return None
    moment = epoch + timedelta(days=int(number))
    return moment.date() if _YEAR_MIN <= moment.year <= _YEAR_MAX else None


def _infer_year(month: int, day: int, today: date | None, prefer_future: bool) -> date | None:
    """Year-less "Sep 20": the most recent such day (verified/posted) or the next upcoming one (deadlines)."""
    if today is None:
        return None
    this_year = _safe_date(today.year, month, day)
    if prefer_future:
        if this_year is not None and this_year >= today - timedelta(days=30):
            return this_year
        return _safe_date(today.year + 1, month, day)
    if this_year is not None and this_year <= today + timedelta(days=7):
        return this_year
    return _safe_date(today.year - 1, month, day)


def _parse_numeric_date(
    text: str, day_first: bool, today: date | None, future: bool
) -> date | None:
    if m := re.match(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})(?:$|[Tt\s,])", text):
        return _safe_date(int(m[1]), int(m[2]), int(m[3]))
    if m := re.fullmatch(r"(\d{4})[-/.](\d{1,2})", text):
        return _month_end(int(m[1]), int(m[2]))
    if m := re.match(r"^(\d{1,2})([/\-.])(\d{1,2})\2(\d{2}|\d{4})(?:$|[Tt\s,])", text):
        a, b, y = int(m[1]), int(m[3]), int(m[4])
        year = _two_digit_year(y) if len(m[4]) == 2 else y
        if m[2] == ".":
            day_first = True  # dotted dates are European: 01.09.2026 is 1 September
        if a > 12 >= b:
            month, day = b, a
        elif b > 12 >= a:
            month, day = a, b
        elif a <= 12 and b <= 12:
            month, day = (b, a) if day_first else (a, b)
        else:
            return None
        return _safe_date(year, month, day)
    if m := re.fullmatch(r"(\d{1,2})[/\-.](\d{1,2})", text):
        a, b = int(m[1]), int(m[2])
        month, day = (b, a) if (day_first or a > 12) else (a, b)
        return _infer_year(month, day, today, future) if 1 <= month <= 12 else None
    return None


def _parse_named_month(lowered: str, today: date | None, future: bool) -> date | None:
    month: int | None = None
    numbers: list[str] = []
    for token in re.findall(r"[a-z]+|\d+", _TIME_OF_DAY.sub(" ", lowered)):
        if token.isdigit():
            numbers.append(token)
        elif token in _MONTHS:
            if month is not None and _MONTHS[token] != month:
                return None
            month = _MONTHS[token]
        elif token not in _WEEKDAYS and token not in _DATE_NOISE:
            return None
    if month is None:
        return None
    years = [n for n in numbers if len(n) == 4]
    smalls = [n for n in numbers if len(n) <= 2]
    if len(years) + len(smalls) != len(numbers) or len(years) > 1:
        return None
    if years:
        if len(smalls) > 1:
            return None
        year = int(years[0])
        return _safe_date(year, month, int(smalls[0])) if smalls else _month_end(year, month)
    if len(smalls) == 2:  # "1-Sep-26": day then two-digit year
        return _safe_date(_two_digit_year(int(smalls[1])), month, int(smalls[0]))
    if len(smalls) == 1:
        return _infer_year(month, int(smalls[0]), today, future)
    return None


def _parse_date_text(text: str, day_first: bool, today: date | None, future: bool) -> date | None:
    low = _LEAD_NOISE.sub("", clean_text(text).strip("()[]").lower()).strip()
    if not low:
        return None
    s = low
    if today is not None:
        if low in ("today", "now"):
            return today
        if low == "yesterday":
            return today - timedelta(days=1)
        if m := _RELATIVE.fullmatch(low):
            return today - timedelta(days=int(m[1]) * _RELATIVE_DAYS[m[2]])
    if re.fullmatch(r"\d{5,8}(?:\.\d+)?", s):
        return _from_number(float(s), _EXCEL_1900)
    if s[0].isdigit() and (parsed := _parse_numeric_date(s, day_first, today, future)):
        return parsed
    return _parse_named_month(low, today, future)


def parse_date_value(
    value: object,
    *,
    day_first: bool = False,
    epoch: datetime | None = None,
    today: date | None = None,
    prefer_future: bool = False,
) -> date | None:
    """Best-effort date from a cell value; ``None`` when it is not (clearly) a date. Never raises.

    Handles real Excel dates, Excel serial numbers (1900 or 1904 epoch), ``YYYYMMDD`` integers and text:
    ``2026-09-01``, ``9/1/26``, ``09/01/2026``, ``01.09.2026`` (dotted = day first), ``Sep 1 2026``,
    ``September 1st, 2026``, ``1 Sep 2026``, ``1-Sep-26``, ``Tue, Sep 1, 2026``, ISO timestamps, "3 days ago".
    Ambiguous ``a/b/yy`` is month-first unless ``day_first`` (or ``a`` > 12). Month + year alone resolves to
    the LAST day of that month (benefit of the doubt for freshness). A year-less "Sep 20" needs ``today``:
    the most recent such day, or with ``prefer_future`` (deadlines) the next upcoming one.
    """
    try:
        if value is None or isinstance(value, bool):
            return None
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        if isinstance(value, dtime | timedelta):
            return None
        if isinstance(value, int | float):
            return _from_number(float(value), epoch or _EXCEL_1900)
        if isinstance(value, str):
            return _parse_date_text(value, day_first, today, prefer_future)
    except (ValueError, OverflowError):
        return None
    return None


_SLASH_DATE = re.compile(r"^\s*(\d{1,2})[/\-](\d{1,2})[/\-](\d{2}|\d{4})\b")


def infer_day_first(values: Iterable[object]) -> bool:
    """True when the text dates of a sheet prove day-first order (``25/09/2026``) and never month-first."""
    day_first = month_first = False
    for value in values:
        if isinstance(value, str) and (m := _SLASH_DATE.match(value)):
            a, b = int(m[1]), int(m[2])
            if a > 12 >= b:
                day_first = True
            elif b > 12 >= a:
                month_first = True
    return day_first and not month_first


# ------------------------------------------------------------------------------------------------ terms


class Term(NamedTuple):
    """A (season, year) mention; either part may be unknown (``"Summer"``, ``"2027"``)."""

    season: str | None
    year: int | None


TermVerdict = Literal["match", "weak", "other", "none"]

_SEASON_WORDS = {
    "summer": "summer",
    "summr": "summer",
    "sum": "summer",
    "fall": "fall",
    "autumn": "fall",
    "spring": "spring",
    "winter": "winter",
}
_SEASON_ABBREVIATIONS = frozenset({"sum", "summr"})  # only trusted next to a year ("sum" is a word)
_SEASON_TOKEN = re.compile(
    r"(?<![A-Za-z])(summer|summr|sum|fall|autumn|spring|winter)(?![A-Za-z])\.?", re.I
)
_YEAR_TOKEN = re.compile(
    r"(?<![\d$#.])(?:(?P<y4>20\d\d)(?!\d)|['’‘`]\s?(?P<y2a>\d\d)(?!\d)|(?P<y2b>\d\d)(?![\d%]))"
)
_UNIT_AFTER = re.compile(r"^\s*-?\s*(?:weeks?|wks?|months?|mos?\b|hours?|hrs?|days?)", re.I)
_SEASON_JOIN = re.compile(r"^\s*(?:/|&|,|-|–|—|\+|and|or)\s*$", re.I)
_GAP_AFTER = re.compile(r"^[\s,.\-–—/']*(?:[A-Za-z]+[\s,.\-–—]*){0,3}$")
_GAP_ADJACENT = re.compile(r"^[\s,.\-–—/']*$")
# "Applications open Fall 2026" / "deadline Sep 2026": a term named as an application window, not the cohort.
_WINDOW_CUE = re.compile(
    r"\b(?:appl(?:y|ies|ication|ications|ying)|deadline|due|opens?|opening|closes?|closing|posted|updated|"
    r"verified|added|recruit(?:ing|ment)?|accepting|as of|since|until|before|by)\b[^.;\n]{0,25}$",
    re.I,
)


class _YearToken(NamedTuple):
    start: int
    end: int
    kind: str  # "y4" 2027 | "y2a" '27 | "y2b" bare 27
    value: int


def _year_tokens(text: str) -> list[_YearToken]:
    return [
        _YearToken(
            m.start(),
            m.end(),
            "y4" if m.group("y4") else ("y2a" if m.group("y2a") else "y2b"),
            int(m.group("y4") or m.group("y2a") or m.group("y2b")),
        )
        for m in _YEAR_TOKEN.finditer(text)
    ]


def _season_groups(text: str) -> list[list[tuple[int, int, str, str]]]:
    """Seasons joined by "/", "&", "or" ("Summer/Fall") form one group sharing a year."""
    groups: list[list[tuple[int, int, str, str]]] = []
    for m in _SEASON_TOKEN.finditer(text):
        item = (m.start(), m.end(), _SEASON_WORDS[m.group(1).lower()], m.group(1).lower())
        if groups and _SEASON_JOIN.match(text[groups[-1][-1][1] : item[0]]):
            groups[-1].append(item)
        else:
            groups.append([item])
    return groups


def _year_fits(text: str, token: _YearToken, gap: str, adjacent_only: bool) -> bool:
    if token.kind == "y2b":  # a bare "27" must hug the season and look like a year, not "12-week"
        return (
            24 <= token.value <= 35
            and not _UNIT_AFTER.match(text[token.end : token.end + 12])
            and bool(_GAP_ADJACENT.match(gap))
        )
    return bool((_GAP_ADJACENT if adjacent_only else _GAP_AFTER).match(gap))


def _scan_terms(text: str) -> list[tuple[Term, int]]:
    """All (season, year) mentions with their start offset; unpaired ones have a ``None`` part."""
    years = _year_tokens(text)
    used: set[int] = set()
    found: list[tuple[Term, int]] = []
    for group in _season_groups(text):
        start, end = group[0][0], group[-1][1]
        chosen: int | None = None
        for idx, tok in enumerate(years):  # year after: "Summer 2027", "Summer Internship 2027"
            if (
                idx not in used
                and tok.start >= end
                and _year_fits(text, tok, text[end : tok.start], False)
            ):
                chosen = idx
                break
        if chosen is None:  # year before: "2027 Summer"
            for idx in range(len(years) - 1, -1, -1):
                tok = years[idx]
                if (
                    idx not in used
                    and tok.end <= start
                    and _year_fits(text, tok, text[tok.end : start], True)
                ):
                    chosen = idx
                    break
        if chosen is not None:
            used.add(chosen)
            tok = years[chosen]
            year = tok.value if tok.kind == "y4" else _two_digit_year(tok.value)
            found.extend((Term(season, year), pos) for pos, _, season, _ in group)
        else:
            found.extend(
                (Term(season, None), pos)
                for pos, _, season, raw in group
                if raw not in _SEASON_ABBREVIATIONS
            )
    found.extend(
        (Term(None, t.value), t.start)
        for i, t in enumerate(years)
        if i not in used and t.kind == "y4"
    )
    return sorted(found, key=lambda pair: pair[1])


def parse_terms(text: str | None, *, strict: bool = False) -> list[Term]:
    """Term mentions in free text: ``Summer 2027``, ``Summer '27``, ``Sum 2027``, ``2027 Summer``, ``Summer/Fall 2027``.

    ``strict=True`` (titles / notes) keeps only full season+year mentions and ignores those phrased as an
    application window ("Applications open Fall 2026"). ``strict=False`` (a Term cell) also returns a bare
    season ("Summer") or a bare year ("2027"). Two-digit years need an apostrophe ("'27") unless they sit next
    to a season and fall in 2024-2035 ("Summer 27" but not "Summer 12-week").
    """
    if not text:
        return []
    body = text[:_SCAN_LIMIT]
    mentions: list[Term] = []
    for term, start in _scan_terms(body):
        if strict:
            if term.season is None or term.year is None:
                continue
            if _WINDOW_CUE.search(body[max(0, start - 40) : start]):
                continue
        mentions.append(term)
    return mentions


@lru_cache(maxsize=64)
def parse_target_term(target: str) -> Term:
    """``search.target_term`` -> ``Term`` (``Term(None, None)`` when it holds neither a season nor a year)."""
    mentions = parse_terms(target)
    return mentions[0] if mentions else Term(None, None)


def _relation(target: Term, mention: Term) -> TermVerdict:
    missing = 0
    for wanted, seen in ((target.season, mention.season), (target.year, mention.year)):
        if wanted is None:
            continue
        if seen is None:
            missing += 1
        elif seen != wanted:
            return "other"
    return "weak" if missing else "match"


def compare_term(target: Term, mentions: Sequence[Term]) -> TermVerdict:
    """Overall verdict: any exact mention -> match; else any partial -> weak; else a different term -> other."""
    if target == Term(None, None) or not mentions:
        return "none"
    relations = {_relation(target, m) for m in mentions}
    if "match" in relations:
        return "match"
    return "weak" if "weak" in relations else "other"


@dataclass(frozen=True)
class TermDecision:
    status: Literal["ok", "assumed", "wrong"]
    detail: str = ""


def decide_term(target_term: str, term_cell: str, title: str, notes: str) -> TermDecision:
    """Does a row belong to ``target_term``?  ``ok`` / ``assumed`` (no evidence either way) / ``wrong``.

    A non-empty Term cell with a recognisable term must match (different season or year -> ``wrong``). An empty
    or unrecognisable cell falls back to explicit season+year mentions in the title / notes; an explicit
    different term there also excludes the row; no evidence at all -> ``assumed`` (keep, flagged).
    """
    target = parse_target_term(target_term)
    if target == Term(None, None):  # unparseable target: plain text comparison
        if not term_cell:
            return TermDecision("assumed")
        wanted = norm_text(target_term)
        if wanted and wanted in norm_text(term_cell):
            return TermDecision("ok")
        return TermDecision("wrong", f"term {term_cell!r} is not {target_term!r}")
    text_verdict = compare_term(
        target, parse_terms(title, strict=True) + parse_terms(notes, strict=True)
    )
    if term_cell:
        cell_verdict = compare_term(target, parse_terms(term_cell))
        if cell_verdict == "match":
            return TermDecision("ok")
        if cell_verdict == "other":
            return TermDecision("wrong", f"term {term_cell!r} is not {target_term!r}")
        if cell_verdict == "weak":  # "Summer" / "2027": the text may settle it, or contradict it
            if text_verdict == "other":
                return TermDecision("wrong", "title/notes name a different term")
            return TermDecision("ok" if text_verdict == "match" else "assumed")
    if text_verdict == "match":
        return TermDecision("ok")
    if text_verdict == "other":
        return TermDecision("wrong", "title/notes name a different term")
    return TermDecision("assumed")


# ------------------------------------------------------------------------------------------------ status

StatusKind = Literal["open", "closed", "applied", "unknown"]

_YES_SYMBOLS = frozenset("✓✔✅☑\U0001f7e2")
_NO_SYMBOLS = frozenset("✗✘✕❌☒✖\U0001f534")
_YES_WORDS = frozenset(_words("yes y true 1 yep yeah"))
_NO_WORDS = frozenset(_words("no n false 0 nope"))
_NOT_CLOSED = re.compile(
    r"\bnot (?:yet |been )?(?:closed|filled|expired)\b|\buntil (?:the )?(?:position |role )?(?:is )?"
    r"(?:filled|closed)\b"
)
_CLOSED_RE = re.compile(
    r"\b(?:closed|filled|expired|inactive|unavailable|cancell?ed|withdrawn|removed|archived|ended|"
    r"deactivated|discontinued|suspended|paused|on hold|taken down|no longer|not accepting|not open|"
    r"not yet open|not hiring|not available|not active|coming soon|opens? (?:soon|later)|upcoming|dead)\b"
)
_APPLIED_RE = re.compile(r"^(?:already )?applied\b|^(?:application )?submitted\b")
_OPEN_RE = re.compile(
    r"\b(?:open|active|verified|live|accepting|available|hiring|rolling|ongoing|current|posted|"
    r"yes|y|true)\b"
)


def classify_status(value: object, *, yes_means: Literal["open", "closed"] = "open") -> StatusKind:
    """Status cell -> ``open`` / ``closed`` / ``applied`` / ``unknown`` (unknown counts as open upstream).

    Open: Open, Active, Verified, Live, Yes/Y/TRUE, a tick. Closed: Closed, Filled, Expired, Inactive, No/N/FALSE,
    a cross, "No longer accepting", "Not open", "On hold". ``yes_means="closed"`` is for "Closed?" / "Filled?"
    columns, where a plain yes/no answer is inverted; explicit words keep their own meaning either way.
    """
    if value is None:
        return "unknown"
    yes: StatusKind = yes_means
    no: StatusKind = "closed" if yes_means == "open" else "open"
    if isinstance(value, bool):
        return yes if value else no
    if isinstance(value, int | float):
        return yes if value == 1 else (no if value == 0 else "unknown")
    text = clean_text(value)
    if not text:
        return "unknown"
    if any(ch in _YES_SYMBOLS for ch in text):
        return yes
    if any(ch in _NO_SYMBOLS for ch in text):
        return no
    n = norm_text(text)
    if not n:
        return "unknown"
    if n in _YES_WORDS:
        return yes
    if n in _NO_WORDS:
        return no
    if _APPLIED_RE.match(n):
        return "applied"
    if _NOT_CLOSED.search(n):
        return "open"
    if _CLOSED_RE.search(n):
        return "closed"
    return "open" if _OPEN_RE.search(n) else "unknown"


_SECTION_CLOSED = re.compile(
    r"\b(?:closed|expired|filled|inactive|archive|archived|past|old|removed)\b"
)
_SECTION_OPEN = re.compile(r"\b(?:open|active|current|live|verified|new)\b")


def section_state(text: str) -> Literal["open", "closed"] | None:
    """A one-cell row such as "Closed / Archived" or "Open roles" switches the state of the rows below it."""
    n = norm_text(text)
    if not n or len(n.split()) > 6:
        return None
    if _SECTION_CLOSED.search(n):
        return "closed"
    return "open" if _SECTION_OPEN.search(n) else None


# ------------------------------------------------------------------------------------------------ internship-ish

_INTERN_RE = re.compile(
    r"\b(?:interns?|internships?|co-?ops?|externships?|apprentice(?:ship)?s?|trainees?|fellowships?|"
    r"summer (?:analysts?|associates?|scholars?|fellows?|programs?|programmes?|students?)|"
    r"student (?:workers?|trainees?|programs?))\b",
    re.I,
)
_NON_INTERN_WORDS = (
    *_words(
        "senior sr staff principal director vp head chief lead executive distinguished ii iii iv fulltime "
        "experienced permanent"
    ),
    "vice president",
    "head of",
    "full time",
    "new grad",
    "new graduate",
    "new grads",
    "entry level",
)


@lru_cache(maxsize=16)
def _non_intern_regex(extra: tuple[str, ...]) -> re.Pattern[str]:
    parts = []
    for word in (*_NON_INTERN_WORDS, *extra):
        n = norm_text(word)
        if n:
            parts.append(r"\s+".join(re.escape(tok) for tok in n.split()))
    return re.compile(r"\b(?:" + "|".join(parts) + r")\b")


def looks_like_internship(
    title: str,
    type_text: str = "",
    notes: str = "",
    *,
    has_term: bool = False,
    extra_non_intern: Iterable[str] = (),
) -> tuple[bool, str]:
    """(is internship-ish, reason). Lenient: only explicit non-intern evidence rejects a row.

    Intern wording in the title (or a Type/Level cell) always accepts. Otherwise senior / staff / director /
    "full-time" / "new grad" style wording in the title or type rejects. Intern wording in the notes, a seasonal
    term, or simply no evidence at all accepts (the sheet is a list of internships).
    """
    if _INTERN_RE.search(title) or _INTERN_RE.search(type_text):
        return True, "intern wording"
    negative = _non_intern_regex(tuple(extra_non_intern))
    if hit := negative.search(f"{norm_text(title)} {norm_text(type_text)}"):
        return False, f"non-internship wording {hit.group(0)!r} without intern wording"
    if _INTERN_RE.search(notes[:_SCAN_LIMIT]):
        return True, "intern wording in notes"
    return True, "seasonal term" if has_term else "no evidence either way"


# ------------------------------------------------------------------------------------------------ header matching

FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    field_name: _words(aliases, "|")
    for field_name, aliases in {
        "company": (
            "company | employer | organization | organisation | company name | employer name | "
            "organization name | organisation name | org | firm | business | corporation | hiring company | "
            "hiring organization | company employer | employer company | company organization | "
            "organization company | companies | employers | brand"
        ),
        "title": (
            "title | role | position | job title | position title | role title | job | job role | "
            "job position | position name | role name | opportunity | opportunity name | internship | "
            "internship title | internship role | internship position | job name | posting title | opening | "
            "job opening | program | program name | intern role | intern position | roles | positions | "
            "job titles | role position | position role | opportunities"
        ),
        "url": (
            "url | link | job link | job url | posting | posting url | posting link | job posting | "
            "job posting url | job posting link | job listing | listing | listing url | listing link | "
            "careers link | careers page | career page | careers url | career link | job page | web link | "
            "website | web address | hyperlink | links | urls | position link | role link | internship link | "
            "internship url | link url | url link | job board link | posting page"
        ),
        "apply_url": (
            "apply link | apply url | application url | application link | apply | apply here | apply now | "
            "apply page | application page | application form | how to apply | apply online | direct link | "
            "direct url | direct apply | direct apply link | application website | apply at | apply via | "
            "link to apply | application links | apply links | application web address | application site | "
            "applications link"
        ),
        "location": (
            "location | city | locations | job location | office | office location | work location | "
            "city state | city and state | primary location | metro | location city | city location | where | "
            "cities | base location | office city"
        ),
        "term": (
            "term | season | internship term | cohort | semester | session | internship season | intern term | "
            "program term | term season | season term | start term | internship cohort | recruiting cycle | "
            "cycle | year | internship year | program year | cohort year | term year | timeframe | time frame | "
            "intern season | hiring term | hiring season"
        ),
        "status": (
            "status | open | state | open closed | open or closed | job status | posting status | "
            "listing status | availability | active | currently open | is open | still open | live | "
            "open status | current status | accepting applications | accepting | applications open | "
            "app status | opportunity status | position status | role status | open y n | status open closed"
        ),
        "closed": (
            "closed | is closed | filled | position filled | expired | is expired | is filled | role filled | "
            "closed y n | job filled"
        ),
        "posted": (
            "posted | date posted | date added | posted date | posting date | date listed | listed | added | "
            "added on | date found | found on | published | publish date | post date | opened | open date | "
            "date opened | date discovered | first seen | date | posted on | listed on | listing date | "
            "created | created on | date created | date of posting | posted at"
        ),
        "verified": (
            "last verified | verified on | date verified | verified | verified date | last checked | "
            "last updated | updated | last update | date updated | checked | date checked | last confirmed | "
            "confirmed on | last seen | verification date | last check | verified at | as of | updated on | "
            "last modified | last verified date | last verified on | date last verified | last verification"
        ),
        "deadline": (
            "deadline | application deadline | apply by | due date | due | closes | close date | "
            "closing date | closing | application due | apply deadline | deadline date | closes on | expires | "
            "expiration | expiry date | expires on | expiration date | apply before | applications close | "
            "app deadline | priority deadline | final deadline | application close date | application closes | "
            "deadline to apply | apply by date | app due | close | expiry | date due | last day to apply"
        ),
        "ats": (
            "ats | platform | application platform | portal | application portal | application system | "
            "ats platform | applicant tracking system | system | portal type | ats portal | ats system | "
            "apply platform | hiring platform | ats type | ats vendor | ats name | portal name | "
            "job platform | career platform"
        ),
        "notes": (
            "notes | note | comments | comment | remarks | remark | additional info | additional information | "
            "info | other | misc | notes comments | comments notes | internal notes | extra notes | "
            "other notes | other info | special notes"
        ),
        "description": (
            "description | job description | details | summary | about | role description | responsibilities | "
            "job summary | overview | position description | job details | role details | "
            "description details | position summary | role summary"
        ),
        "type": (
            "type | job type | employment type | level | job level | position type | role type | "
            "opportunity type | internship type | employment | kind | classification | role level | "
            "position level"
        ),
    }.items()
}
FIELDS: tuple[str, ...] = tuple(FIELD_ALIASES)

_ALIAS_LOOKUP: dict[str, tuple[str, int]] = {}
_ALIAS_COMPACT: dict[str, tuple[str, int]] = {}
for _field_name, _aliases in FIELD_ALIASES.items():
    for _rank, _alias in enumerate(_aliases):
        _key = norm_text(_alias)
        _ALIAS_LOOKUP.setdefault(_key, (_field_name, _rank))
        _ALIAS_COMPACT.setdefault(_key.replace(" ", ""), (_field_name, _rank))
_ALIAS_WORDS = frozenset(w for key in _ALIAS_LOOKUP for w in key.split())
_FUZZY_WORDS = sorted(w for w in _ALIAS_WORDS if len(w) >= 5)
_HEADER_NOISE = frozenset(
    _words("the of required optional if any mm dd yy yyyy est cst pst utc format")
)
_PARENTHETICAL = re.compile(r"[(\[{][^)\]}]*[)\]}]")


class HeaderMatch(NamedTuple):
    field: str
    score: float  # 2.0 column_map override, 1.0 exact alias, 0.95 spaceless, 0.8 covered, 0.7 fuzzy
    rank: int  # position of the alias in its table: earlier = preferred among duplicate columns


def base_field(name: str) -> str:
    """``apply_url`` counts as ``url`` when deciding whether a row looks like a header."""
    return "url" if name == "apply_url" else name


def _cover(
    tokens: tuple[str, ...], name: str | None = None, best: int = 99
) -> tuple[str, int] | None:
    """(field, best rank) when consecutive known alias phrases of ONE field account for every token."""
    if not tokens:
        return (name, best) if name else None
    for size in range(min(3, len(tokens)), 0, -1):
        hit = _ALIAS_LOOKUP.get(" ".join(tokens[:size]))
        if (
            hit
            and name in (None, hit[0])
            and (rest := _cover(tokens[size:], hit[0], min(best, hit[1])))
        ):
            return rest
    return None


def _correct_word(word: str) -> str:
    """Nearest known header word for a misspelt one (>= 5 letters, unambiguous), else the word itself."""
    if word in _ALIAS_WORDS or len(word) < 5:
        return word
    close = difflib.get_close_matches(word, _FUZZY_WORDS, n=1, cutoff=0.8)
    return close[0] if close else word


@lru_cache(maxsize=4096)
def _match_header_text(text: str, fuzzy: bool) -> HeaderMatch | None:
    stripped = norm_text(_PARENTHETICAL.sub(" ", text))
    for candidate in dict.fromkeys(c for c in (stripped, norm_text(text)) if c):
        if hit := _ALIAS_LOOKUP.get(candidate):
            return HeaderMatch(hit[0], 1.0, hit[1])
        if hit := _ALIAS_COMPACT.get(candidate.replace(" ", "")):
            return HeaderMatch(hit[0], 0.95, hit[1])
    tokens = tuple(t for t in stripped.split() if t not in _HEADER_NOISE)
    if 1 < len(tokens) <= 4 and (covered := _cover(tokens)):
        return HeaderMatch(covered[0], 0.8, covered[1])
    if fuzzy and 1 <= len(tokens) <= 3:  # a typo in one word: "Compnay", "Aplly Link", "Postion"
        fixed = tuple(_correct_word(t) for t in tokens)
        if fixed != tokens:
            if hit := _ALIAS_LOOKUP.get(" ".join(fixed)):
                return HeaderMatch(hit[0], 0.7, hit[1])
            if len(fixed) > 1 and (covered := _cover(fixed)):
                return HeaderMatch(covered[0], 0.7, covered[1])
    return None


def match_header(text: object, *, fuzzy: bool = True) -> HeaderMatch | None:
    """Recognise a column header (case, spacing, punctuation insensitive; trailing ``?`` / ``:`` tolerated).

    Order: exact alias (parentheticals such as "(required)" are ignored), spaceless alias (``JobTitle``),
    several aliases of ONE field that together account for every word (``"Employer / Organization"``), then a
    close typo (``"Compnay"``). Anything else - ``"Company Size"``, ``"Company Location"`` - is not matched.
    """
    if not isinstance(text, str) or len(text) > 80:
        return None
    return _match_header_text(text, fuzzy)


def canonical_field(name: str) -> str | None:
    """``column_map`` key -> canonical field ("role" -> "title", "apply_url" -> "apply_url"); else None."""
    n = norm_text(name)
    if n.replace(" ", "_") in FIELD_ALIASES:
        return n.replace(" ", "_")
    hit = _ALIAS_LOOKUP.get(n)
    return hit[0] if hit else None


# ------------------------------------------------------------------------------------------------ sheet grid


@dataclass(frozen=True, slots=True)
class Cell:
    """One cell: its cached value and (when known) the external hyperlink it carries."""

    value: Any = None
    link: str | None = None


EMPTY_CELL = Cell()


@dataclass
class SheetGrid:
    """A worksheet as a plain grid. ``rows[0]`` is sheet row 1; every row has the same width."""

    name: str
    rows: list[list[Cell]]
    epoch: datetime = _EXCEL_1900
    state: str = "visible"


@dataclass(frozen=True)
class HeaderInfo:
    """Where the header is and which column holds what."""

    row: int  # 1-based sheet row of the header
    headers: tuple[str, ...]  # cleaned header text per column ("" when blank)
    fields: dict[str, tuple[int, ...]]  # canonical field -> 0-based column indexes, best first
    extras: dict[int, str]  # unmapped column -> unique key used in ``Opportunity.extra``
    warnings: tuple[str, ...] = ()

    def primary_headers(self) -> dict[str, str]:
        return {f: self.headers[cols[0]] for f, cols in self.fields.items() if cols}


class _Overrides(NamedTuple):
    by_header: dict[str, str]  # spaceless normalised header text -> field
    by_letter: dict[str, str]  # column letters ("C") -> field
    notes: list[str]


def _override_index(column_map: Mapping[str, str]) -> _Overrides:
    by_header: dict[str, str] = {}
    by_letter: dict[str, str] = {}
    notes: list[str] = []
    for key, header in column_map.items():
        name = canonical_field(key)
        if name is None:
            notes.append(f"column_map: unknown field {key!r} (ignored)")
            continue
        text = clean_text(header)
        if not text:
            continue
        by_header[norm_text(text).replace(" ", "")] = name
        if re.fullmatch(r"[A-Za-z]{1,3}", text):
            by_letter[text.upper()] = name
    return _Overrides(by_header, by_letter, notes)


def _match_row(
    cells: Sequence[Cell], overrides: _Overrides | None = None, *, fuzzy: bool = True
) -> dict[int, HeaderMatch]:
    matches: dict[int, HeaderMatch] = {}
    for col, cell in enumerate(cells):
        if not isinstance(cell.value, str):
            continue
        name = (
            overrides.by_header.get(norm_text(cell.value).replace(" ", "")) if overrides else None
        )
        if name:
            matches[col] = HeaderMatch(name, 2.0, 0)
        elif (hit := match_header(cell.value, fuzzy=fuzzy)) is not None:
            matches[col] = hit
    if overrides:
        for letter, name in overrides.by_letter.items():  # an override given as a column letter
            col = column_index_from_string(letter) - 1
            satisfied = any(m.score >= 2.0 and m.field == name for m in matches.values())
            if col < len(cells) and col not in matches and not satisfied:
                matches[col] = HeaderMatch(name, 2.0, 0)
    return matches


def detect_header(
    rows: Sequence[Sequence[Cell]],
    column_map: Mapping[str, str] | None = None,
    *,
    scan: int = HEADER_SCAN_ROWS,
) -> HeaderInfo | None:
    """First row within ``scan`` rows with >= 3 recognised columns (distinct fields), or None.

    ``column_map`` (canonical field -> header text or column letter) overrides win over the alias tables and
    count as recognised. Duplicate headers keep every column (best first); unmapped headers become extras.
    """
    overrides = _override_index(column_map or {})
    for idx, cells in enumerate(rows[:scan]):
        matches = _match_row(cells, overrides)
        if len({base_field(m.field) for m in matches.values()}) >= MIN_HEADER_FIELDS:
            return _build_header(idx + 1, cells, matches, overrides)
    return None


def _build_header(
    row_no: int, cells: Sequence[Cell], matches: dict[int, HeaderMatch], overrides: _Overrides
) -> HeaderInfo:
    headers = tuple(clean_text(c.value, limit=120) for c in cells)
    notes = list(overrides.notes)
    matched = {h.field for h in matches.values() if h.score >= 2.0}
    notes.extend(
        f"column_map: no column matches the override for {name!r}"
        for name in dict.fromkeys(overrides.by_header.values())
        if name not in matched
    )
    by_field: dict[str, list[tuple[float, int, int]]] = {}
    for col, m in matches.items():
        by_field.setdefault(m.field, []).append((-m.score, m.rank, col))
    fields = {name: tuple(col for _, _, col in sorted(items)) for name, items in by_field.items()}
    extras: dict[int, str] = {}
    seen: dict[str, int] = {}
    for col, text in enumerate(headers):
        if not text or col in matches:
            continue
        seen[text] = seen.get(text, 0) + 1
        extras[col] = text if seen[text] == 1 else f"{text} ({seen[text]})"
    return HeaderInfo(row_no, headers, fields, extras, tuple(notes))


# ------------------------------------------------------------------------------------------------ reading files

_HYPERLINK_CALL = re.compile(r"^\s*=?\s*HYPERLINK\s*\((.*)\)\s*$", re.I | re.S)
_A1_REF = re.compile(r"^\$?([A-Za-z]{1,3})\$?(\d{1,7})$")


def _split_args(body: str) -> list[str]:
    """Split a function's argument text at top-level commas / semicolons (quotes and parentheses respected)."""
    args: list[str] = []
    current: list[str] = []
    depth = 0
    in_quote = False
    i = 0
    while i < len(body):
        ch = body[i]
        if in_quote:
            current.append(ch)
            if ch == '"':
                if body[i + 1 : i + 2] == '"':  # "" is an escaped quote inside a string
                    current.append('"')
                    i += 1
                else:
                    in_quote = False
        elif ch in ",;" and depth == 0:
            args.append("".join(current).strip())
            current = []
        else:
            in_quote = ch == '"'
            depth += (ch == "(") - (ch == ")")
            current.append(ch)
        i += 1
    args.append("".join(current).strip())
    return args


def _eval_string_expr(expr: str, lookup: Callable[[int, int], object]) -> str | None:
    """Tiny evaluator: string literals, same-sheet cell references and ``&`` concatenation; else None."""
    pieces: list[str] = []
    part: list[str] = []
    in_quote = False
    for ch in expr:
        if ch == '"':
            in_quote = not in_quote
        if ch == "&" and not in_quote:
            pieces.append("".join(part).strip())
            part = []
        else:
            part.append(ch)
    pieces.append("".join(part).strip())
    out: list[str] = []
    for piece in pieces:
        if len(piece) >= 2 and piece.startswith('"') and piece.endswith('"'):
            out.append(piece[1:-1].replace('""', '"'))
        elif m := _A1_REF.match(piece):
            value = lookup(int(m[2]), column_index_from_string(m[1].upper()))
            if value is None:
                return None
            out.append(clean_text(value))
        else:
            return None
    return "".join(out)


def parse_hyperlink_formula(
    formula: str, lookup: Callable[[int, int], object] = lambda row, col: None
) -> tuple[str | None, str | None] | None:
    """``=HYPERLINK("https://..","Apply")`` -> (url, label). ``None`` if it is not a HYPERLINK formula.

    Arguments may be literals, ``A1`` references (resolved through ``lookup(row, col)``) or ``&`` chains of
    both; an argument that cannot be evaluated yields None for that half.
    """
    match = _HYPERLINK_CALL.match(formula)
    if not match:
        return None
    args = _split_args(match.group(1))
    if not args or not args[0]:
        return None
    target = _eval_string_expr(args[0], lookup)
    label = _eval_string_expr(args[1], lookup) if len(args) > 1 else None
    return (extract_url(target) if target else None), label


def _make_cell(value: Any, hyperlink: Any, data_type: str) -> Cell:
    if data_type == "e" or isinstance(value, dtime | timedelta):
        value = None
    target = getattr(hyperlink, "target", None) if hyperlink is not None else None
    link = str(target) if target else None
    return Cell(value, link) if (value is not None or link) else EMPTY_CELL


def _vertical_merges(ws: Any) -> dict[tuple[int, int], Cell]:
    """Values of vertically merged single-column ranges (a company spanning several role rows)."""
    fills: dict[tuple[int, int], Cell] = {}
    for rng in getattr(getattr(ws, "merged_cells", None), "ranges", ()):
        if rng.min_col != rng.max_col or rng.max_row <= rng.min_row:
            continue
        top = ws.cell(rng.min_row, rng.min_col)
        base = _make_cell(top.value, getattr(top, "hyperlink", None), getattr(top, "data_type", ""))
        if base is EMPTY_CELL:
            continue
        for row in range(rng.min_row + 1, min(rng.max_row, MAX_ROWS) + 1):
            fills[(row, rng.min_col)] = base
    return fills


def _pad_and_trim(rows: list[list[Cell]]) -> list[list[Cell]]:
    while rows and all(c is EMPTY_CELL for c in rows[-1]):
        rows.pop()
    width = max((len(r) for r in rows), default=0)
    for r in rows:
        r.extend([EMPTY_CELL] * (width - len(r)))
    return rows


def _read_rows(ws: Any, *, streaming: bool) -> list[list[Cell]]:
    rows: list[list[Cell]] = []
    if streaming:
        ws.reset_dimensions()  # some writers record a bogus <dimension>
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if row_idx > MAX_ROWS:
                break
            rows.append([_make_cell(v, None, "") for v in row[:MAX_COLS]])
        return _pad_and_trim(rows)
    n_rows = min(ws.max_row or 0, MAX_ROWS)
    n_cols = min(ws.max_column or 0, MAX_COLS)
    if not n_rows or not n_cols:
        return []
    merged = _vertical_merges(ws)
    for r_idx, row in enumerate(
        ws.iter_rows(min_row=1, max_row=n_rows, min_col=1, max_col=n_cols), 1
    ):
        out: list[Cell] = []
        for c_idx, cell in enumerate(row, start=1):
            fill = merged.get((r_idx, c_idx))
            out.append(
                fill
                if fill is not None
                else _make_cell(
                    cell.value, getattr(cell, "hyperlink", None), getattr(cell, "data_type", "")
                )
            )
        rows.append(out)
    return _pad_and_trim(rows)


def _apply_hyperlink_formulas(
    rows: list[list[Cell]], formulas: Mapping[tuple[int, int], str]
) -> None:
    def lookup(row: int, col: int) -> object:
        return (
            rows[row - 1][col - 1].value
            if 1 <= row <= len(rows) and 1 <= col <= len(rows[0])
            else None
        )

    for (row, col), formula in formulas.items():
        if not (1 <= row <= len(rows) and 1 <= col <= len(rows[0])):
            continue
        parsed = parse_hyperlink_formula(formula, lookup)
        if parsed is None:
            continue
        url, label = parsed
        cell = rows[row - 1][col - 1]
        rows[row - 1][col - 1] = Cell(
            cell.value if cell.value not in (None, "") else label, url or cell.link
        )


class _Book:
    """An opened workbook: cached values (+ hyperlinks + merged ranges) and ``=HYPERLINK()`` formulas."""

    def __init__(self, data: bytes, *, streaming: bool) -> None:
        self._data = data
        self.streaming = streaming
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            self._wb = load_workbook(
                io.BytesIO(data), read_only=streaming, data_only=True, keep_links=False
            )
        self._grids: dict[str, SheetGrid] = {}

    @property
    def sheets(self) -> list[tuple[str, str]]:
        """(name, state) of every worksheet, in workbook order."""
        return [
            (ws.title, str(getattr(ws, "sheet_state", "visible"))) for ws in self._wb.worksheets
        ]

    def _hyperlink_formulas(self, name: str) -> dict[tuple[int, int], str]:
        found: dict[tuple[int, int], str] = {}
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                book = load_workbook(
                    io.BytesIO(self._data), read_only=True, data_only=False, keep_links=False
                )
            try:
                ws = book[name]
                ws.reset_dimensions()
                for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                    if r_idx > MAX_ROWS:
                        break
                    for c_idx, value in enumerate(row[:MAX_COLS], start=1):
                        if (
                            isinstance(value, str)
                            and value.startswith("=")
                            and "HYPERLINK" in value.upper()
                        ):
                            found[(r_idx, c_idx)] = value
            finally:
                book.close()
        except Exception as exc:  # formulas are a bonus; the cached values are still usable
            _LOG.debug("workbook: could not read formulas of %r: %s", name, exc)
        return found

    def grid(self, name: str) -> SheetGrid:
        if name not in self._grids:
            ws = self._wb[name]
            rows = _read_rows(ws, streaming=self.streaming)
            if formulas := self._hyperlink_formulas(name):
                _apply_hyperlink_formulas(rows, formulas)
            epoch = getattr(self._wb, "epoch", _EXCEL_1900)
            self._grids[name] = SheetGrid(
                name, rows, epoch, str(getattr(ws, "sheet_state", "visible"))
            )
        return self._grids[name]

    def close(self) -> None:
        try:
            self._wb.close()
        except Exception:  # closing is best effort
            _LOG.debug("workbook: close failed", exc_info=True)


def _open_book(path: Path, *, streaming: bool | None = None) -> _Book:
    """Read the file into memory (no lingering handle; works while Excel has it open) and open it."""
    if path.suffix.lower() == ".xls":
        raise WorkbookError(
            f"{path.name}: the old .xls format is not supported; save it as .xlsx first"
        )
    try:
        data = path.read_bytes()
    except FileNotFoundError as exc:
        raise WorkbookError(f"workbook not found: {path}") from exc
    except PermissionError as exc:
        raise WorkbookError(f"cannot read {path} (close it in Excel and retry): {exc}") from exc
    except OSError as exc:
        raise WorkbookError(f"cannot read {path}: {exc}") from exc
    if not data:
        raise WorkbookError(f"{path.name} is empty")
    stream = len(data) > LARGE_FILE_BYTES if streaming is None else streaming
    try:
        return _Book(data, streaming=stream)
    except Exception as exc:
        raise WorkbookError(
            f"{path.name} is not a readable .xlsx workbook ({type(exc).__name__}: {exc})"
        ) from exc


# ------------------------------------------------------------------------------------------------ sheet selection

WANTED_SHEET = "verified opportunities"
_SHEET_NAME_MIN = 50.0
_SHEET_NEGATIVE = frozenset(
    _words(
        "rejected reject unverified archive archived closed expired old backup junk summary notes stats "
        "dashboard template instructions readme log tmp temp copy"
    )
)


def _stem(token: str) -> str:
    if token in ("opp", "opps"):
        return "opportunity"
    if token.endswith("ies") and len(token) > 4:
        return token[:-3] + "y"
    return token[:-1] if token.endswith("s") and len(token) > 3 else token


def score_sheet_name(name: str, wanted: str = WANTED_SHEET) -> float:
    """0..100 similarity of a sheet name to ``wanted`` (case / spacing / punctuation insensitive).

    ``Verified-Opportunities``, ``verified_opportunities`` and ``Verified Opps`` score >= 90; a lone
    "Opportunities" ~55; "Rejected Opportunities" is penalised; unrelated names score by string similarity.
    """
    n, w = norm_text(name), norm_text(wanted)
    if not n or not w:
        return 0.0
    if n == w:
        return 100.0
    if n.replace(" ", "") == w.replace(" ", ""):
        return 98.0
    name_tokens = set(n.split())
    stems, wanted_stems = {_stem(t) for t in name_tokens}, {_stem(t) for t in w.split()}
    score = difflib.SequenceMatcher(None, n.replace(" ", ""), w.replace(" ", "")).ratio() * 80
    if wanted_stems <= stems:
        score = max(score, 90.0)
    elif wanted == WANTED_SHEET and "opportunity" in stems:
        score = max(score, 55.0)
    elif wanted == WANTED_SHEET and "verified" in stems:
        score = max(score, 50.0)
    if name_tokens & _SHEET_NEGATIVE and not (set(w.split()) & _SHEET_NEGATIVE):
        score -= 40.0
    return max(0.0, min(100.0, score))


def _match_sheet_name(wanted: str, names: Sequence[str]) -> str | None:
    for name in names:
        if name == wanted or name.strip().lower() == wanted.strip().lower():
            return name
    best = max(names, key=lambda n: score_sheet_name(n, wanted), default=None)
    # only spelling differences (case, spacing, punctuation) count: a partial name could pick the wrong tab
    return best if best is not None and score_sheet_name(best, wanted) >= 98.0 else None


@dataclass
class _Selection:
    grid: SheetGrid
    header: HeaderInfo


def _no_header_message(sheet: str | None, names: Sequence[str] = ()) -> str:
    where = f"sheet {sheet!r}" if sheet else f"any sheet ({', '.join(names)})"
    return (
        f"no header row found in {where}: expected a row within the first {HEADER_SCAN_ROWS} rows with at "
        f"least {MIN_HEADER_FIELDS} of Company / Role / Link / Location / Term / Status / Last Verified ... "
        "(set workbook.column_map to map custom headers)"
    )


def _select_sheet(book: _Book, wanted: str | None, column_map: Mapping[str, str]) -> _Selection:
    """``workbook.sheet`` if set, else the visible sheet best matching "verified opportunities" that has a
    detectable header, else (names tell nothing) the sheet with the richest header and the most rows."""
    sheets = book.sheets
    if not sheets:
        raise WorkbookError("the workbook has no worksheets")
    names = [n for n, _ in sheets]
    if wanted and wanted.strip():
        name = _match_sheet_name(wanted.strip(), names)
        if name is None:
            raise WorkbookError(f"sheet {wanted!r} not found; sheets are: {', '.join(names)}")
        header = detect_header(book.grid(name).rows, column_map)
        if header is None:
            raise WorkbookError(_no_header_message(name))
        return _Selection(book.grid(name), header)
    order = sorted(
        range(len(sheets)),
        key=lambda i: (sheets[i][1] != "visible", -score_sheet_name(sheets[i][0]), i),
    )
    named = [
        i
        for i in order
        if sheets[i][1] == "visible" and score_sheet_name(sheets[i][0]) >= _SHEET_NAME_MIN
    ]
    for i in named:
        grid = book.grid(sheets[i][0])
        if (header := detect_header(grid.rows, column_map)) is not None:
            return _Selection(grid, header)
    best: _Selection | None = None
    best_key = (-1, -1)
    for i in (j for j in order if j not in named):
        grid = book.grid(sheets[i][0])
        if (header := detect_header(grid.rows, column_map)) is None:
            continue
        key = (len({base_field(f) for f in header.fields}), len(grid.rows) - header.row)
        if key > best_key:
            best, best_key = _Selection(grid, header), key
    if best is None:
        raise WorkbookError(_no_header_message(None, names))
    return best


# ------------------------------------------------------------------------------------------------ row parsing


@dataclass(frozen=True)
class Rejection:
    """A data row that did not become an opportunity."""

    row: int  # 1-based sheet row
    company: str
    title: str
    reason: RejectReason
    detail: str = ""
    also: tuple[RejectReason, ...] = ()  # other filters the row would have failed as well

    @property
    def label(self) -> str:
        """How summaries name the row: its title, else its company, else "row N"."""
        return self.title or self.company or f"row {self.row}"


@dataclass
class WorkbookParseResult:
    """Outcome of reading one sheet: kept opportunities plus a reason for every rejected row."""

    sheet: str
    header_row: int
    mapping: dict[str, str]
    opportunities: list[Opportunity] = field(default_factory=list)
    rejections: list[Rejection] = field(default_factory=list)
    data_rows: int = 0  # rows with content (kept + rejected)
    junk_rows: int = 0  # blank / separator / repeated-header rows (ignored)
    warnings: list[str] = field(default_factory=list)

    def rejection_summary(self) -> dict[str, list[str]]:
        """``{reason: [row titles]}`` for every rejected row (a title-less row falls back to its company)."""
        out: dict[str, list[str]] = {}
        for rej in self.rejections:
            out.setdefault(rej.reason.value, []).append(rej.label)
        return out

    def rejection_counts(self) -> dict[str, int]:
        return {reason: len(labels) for reason, labels in self.rejection_summary().items()}


def _same_url(a: str, b: str) -> bool:
    return canonical_url(a) == canonical_url(b)


@dataclass
class _Parser:
    header: HeaderInfo
    config: AppConfig
    today: date
    epoch: datetime
    day_first: bool
    sheet: str
    section: Literal["open", "closed"] | None = None
    section_label: str = ""

    # -- cell access ------------------------------------------------------------------------------
    def cells_of(self, cells: Sequence[Cell], name: str) -> list[Cell]:
        return [cells[c] for c in self.header.fields.get(name, ()) if c < len(cells)]

    def text(self, cells: Sequence[Cell], name: str, *, limit: int = _TEXT_LIMIT) -> str:
        for cell in self.cells_of(cells, name):
            if text := clean_text(cell.value, limit=limit):
                return text
        return ""

    def texts(self, cells: Sequence[Cell], name: str, *, multiline: bool = False) -> list[str]:
        out = [clean_text(c.value, multiline=multiline) for c in self.cells_of(cells, name)]
        return [t for t in out if t]

    def date_of(
        self, cells: Sequence[Cell], name: str, *, prefer_future: bool = False
    ) -> date | None:
        for cell in self.cells_of(cells, name):
            parsed = parse_date_value(
                cell.value,
                day_first=self.day_first,
                epoch=self.epoch,
                today=self.today,
                prefer_future=prefer_future,
            )
            if parsed is not None:
                return parsed
        return None

    def url_of(self, cells: Sequence[Cell], name: str) -> str | None:
        for cell in self.cells_of(cells, name):  # a hyperlink target beats the display text
            for candidate in (cell.link, cell.value if isinstance(cell.value, str) else None):
                if candidate and (url := extract_url(candidate)):
                    return url
        return None

    def pick_urls(self, cells: Sequence[Cell]) -> tuple[str | None, str | None]:
        """(url, apply_url): posting and apply columns stay apart; missing both, fall back to stray links."""
        posting = self.url_of(cells, "url")
        apply = self.url_of(cells, "apply_url")
        if not posting and not apply:
            apply = self.url_of(cells, "ats")  # "Application Portal" columns often hold links
        if not posting and not apply:
            # A title cell that is itself a hyperlink, then any link in an unrecognised column. Never the
            # company cell: it usually points at the employer's home page.
            stray = self.cells_of(cells, "title")
            stray += [cells[col] for col in self.header.extras if col < len(cells)]
            for cell in stray:
                if cell.link and (url := extract_url(cell.link)):
                    return url, None
            return None, None
        if posting and apply and _same_url(posting, apply):
            apply = None
        return (posting, apply) if posting else (apply, None)

    def note_separator(self, text: str) -> None:
        if (state := section_state(text)) is not None:
            self.section = state
        self.section_label = clean_text(text, limit=120)


def _is_header_like(cells: Sequence[Cell]) -> bool:
    matches = _match_row(cells, fuzzy=False)
    return len({base_field(m.field) for m in matches.values()}) >= MIN_HEADER_FIELDS


def _status_kind(p: _Parser, cells: Sequence[Cell]) -> StatusKind:
    kinds = [classify_status(c.value) for c in p.cells_of(cells, "status")]
    kinds += [classify_status(c.value, yes_means="closed") for c in p.cells_of(cells, "closed")]
    if "closed" in kinds:
        return "closed"
    if "applied" in kinds:
        return "applied"
    if "open" in kinds:
        return "open"
    return "closed" if p.section == "closed" else "unknown"


def _parse_row(p: _Parser, row_no: int, cells: Sequence[Cell]) -> Opportunity | Rejection | None:
    """One data row -> Opportunity, Rejection, or None for blank / separator / repeated-header rows."""
    filled = [c for c in cells if c.link or clean_text(c.value)]
    if not filled:
        return None
    if len(filled) == 1 or _is_header_like(cells):
        if len(filled) == 1 and (text := clean_text(filled[0].value)):
            p.note_separator(text)
        return None

    search = p.config.search
    company = p.text(cells, "company", limit=_COMPANY_LIMIT)
    title = p.text(cells, "title", limit=_TITLE_LIMIT)
    if not company or not title:
        return Rejection(
            row_no,
            company,
            title,
            RejectReason.MISSING_FIELDS,
            "no company" if not company else "no title",
        )

    reasons: list[tuple[RejectReason, str]] = []
    status = _status_kind(p, cells)
    if status == "closed":
        reasons.append(
            (RejectReason.CLOSED, p.text(cells, "status") or f"in section {p.section_label!r}")
        )
    elif status == "applied":
        reasons.append((RejectReason.ALREADY_APPLIED, p.text(cells, "status")))

    deadline = p.date_of(cells, "deadline", prefer_future=True)
    if deadline is not None and deadline < p.today:
        reasons.append((RejectReason.DEADLINE_PASSED, f"deadline {deadline.isoformat()}"))

    term_cell = p.text(cells, "term")
    notes = " ".join(p.texts(cells, "notes"))
    description = "\n\n".join(p.texts(cells, "description", multiline=True))
    free_text = f"{notes} {description[:_SCAN_LIMIT]}".strip()
    decision = decide_term(search.target_term, term_cell, title, free_text)
    if decision.status == "wrong":
        reasons.append((RejectReason.WRONG_TERM, decision.detail))

    is_intern, why = looks_like_internship(
        title,
        p.text(cells, "type"),
        free_text,
        has_term=bool(term_cell),
        extra_non_intern=search.exclude_title_keywords,
    )
    if not is_intern:
        reasons.append((RejectReason.NOT_INTERNSHIP, why))

    verified = p.date_of(cells, "verified")
    posted = p.date_of(cells, "posted")
    reference = max((d for d in (verified, posted) if d is not None), default=None)
    if reference is not None and (p.today - reference).days > search.recent_days:
        reasons.append(
            (RejectReason.STALE, f"{reference.isoformat()} is older than {search.recent_days} days")
        )

    url, apply_url = p.pick_urls(cells)
    if url is None:
        reasons.append((RejectReason.NO_URL, "no usable link"))

    if reasons or url is None:
        primary, detail = reasons[0]
        return Rejection(row_no, company, title, primary, detail, tuple(r for r, _ in reasons[1:]))
    return _build_opportunity(
        p,
        row_no,
        cells,
        company=company,
        title=title,
        url=url,
        apply_url=apply_url,
        description=description or notes,
        notes=notes,
        term_cell=term_cell,
        term_assumed=decision.status == "assumed",
        dates=(posted, verified, deadline),
        reference=reference,
    )


_RESERVED_EXTRA = frozenset(
    _words(
        "sheet sheet_row section term_assumed date_unknown term_raw status_raw notes posted_raw "
        "last_verified_raw deadline_raw type ats_raw"
    )
)


def _build_opportunity(
    p: _Parser,
    row_no: int,
    cells: Sequence[Cell],
    *,
    company: str,
    title: str,
    url: str,
    apply_url: str | None,
    description: str,
    notes: str,
    term_cell: str,
    term_assumed: bool,
    dates: tuple[date | None, date | None, date | None],
    reference: date | None,
) -> Opportunity:
    posted, verified, deadline = dates
    extra: dict[str, Any] = {"sheet": p.sheet, "sheet_row": row_no}
    if p.section_label:
        extra["section"] = p.section_label
    if term_assumed:
        extra["term_assumed"] = True
    if reference is None:
        extra["date_unknown"] = True
    if term_cell:
        extra["term_raw"] = term_cell
    if status_text := p.text(cells, "status"):
        extra["status_raw"] = status_text
    if notes:
        extra["notes"] = clean_text(notes, limit=_TEXT_LIMIT)
    for parsed, source_field, key in (
        (posted, "posted", "posted_raw"),
        (verified, "verified", "last_verified_raw"),
        (deadline, "deadline", "deadline_raw"),
    ):
        if parsed is None and (raw := p.text(cells, source_field)):
            extra[key] = raw
    if type_text := p.text(cells, "type"):
        extra["type"] = type_text
    ats_text = p.text(cells, "ats")
    ats_hint = parse_ats_hint(ats_text) if ats_text and extract_url(ats_text) is None else None
    if ats_text and ats_hint is None and extract_url(ats_text) is None:
        extra["ats_raw"] = ats_text
    for col, key in p.header.extras.items():
        value = clean_text(cells[col].value, limit=_TEXT_LIMIT) if col < len(cells) else ""
        if value:
            extra[f"{key} (column)" if key in _RESERVED_EXTRA else key] = value

    ats = detect_ats(apply_url or url)
    if ats in (ATS.UNKNOWN, ATS.CUSTOM) and ats_hint is not None:
        ats = (
            ats_hint  # the URL alone is not decisive (employer page that redirects to Workday, ...)
        )
    return Opportunity(
        company=company,
        title=title,
        url=url,
        apply_url=apply_url,
        location=p.text(cells, "location") or None,
        term=p.config.search.target_term,
        source=OpportunitySource.WORKBOOK,
        ats=ats,
        is_open=True,
        posted_date=posted,
        last_verified=verified,
        deadline=deadline,
        description=clean_text(description, multiline=True, limit=_DESCRIPTION_LIMIT) or None,
        extra=extra,
    )


def parse_sheet(
    grid: SheetGrid,
    header: HeaderInfo,
    config: AppConfig,
    *,
    today: date,
    log: logging.Logger | None = None,
) -> WorkbookParseResult:
    """Turn the rows below ``header`` into opportunities + rejections. Never raises for a bad row."""
    logger = log or _LOG
    date_columns = [
        c for name in ("posted", "verified", "deadline") for c in header.fields.get(name, ())
    ]
    day_first = infer_day_first(
        row[c].value for row in grid.rows[header.row :] for c in date_columns if c < len(row)
    )
    parser = _Parser(header, config, today, grid.epoch, day_first, grid.name)
    result = WorkbookParseResult(
        grid.name, header.row, header.primary_headers(), warnings=list(header.warnings)
    )
    if len(grid.rows) >= MAX_ROWS:
        result.warnings.append(f"only the first {MAX_ROWS} rows were read")
    for row_no, cells in enumerate(grid.rows[header.row :], start=header.row + 1):
        try:
            outcome = _parse_row(parser, row_no, cells)
        except Exception as exc:  # a bad row must never abort the ingest
            logger.warning(
                "workbook: skipping row %d of %r: %s: %s",
                row_no,
                grid.name,
                type(exc).__name__,
                exc,
            )
            outcome = Rejection(row_no, "", "", RejectReason.ERROR, f"{type(exc).__name__}: {exc}")
        if outcome is None:
            result.junk_rows += 1
        elif isinstance(outcome, Rejection):
            result.data_rows += 1
            result.rejections.append(outcome)
            logger.debug(
                "workbook: row %d (%s) rejected: %s %s",
                row_no,
                outcome.label,
                outcome.reason,
                outcome.detail,
            )
        else:
            result.data_rows += 1
            result.opportunities.append(outcome)
    logger.info(
        "workbook: %r kept %d of %d rows (%s)",
        grid.name,
        len(result.opportunities),
        result.data_rows,
        ", ".join(f"{k}={v}" for k, v in sorted(result.rejection_counts().items()))
        or "nothing rejected",
    )
    return result


# ------------------------------------------------------------------------------------------------ public API


def resolve_workbook_path(
    raw: str | os.PathLike[str] | None, paths: AppPaths | None = None
) -> Path:
    """``config.workbook.path`` -> a ``Path``: strips quotes ("Copy as path"), expands ``~`` and ``%VAR%``.

    A relative path is tried against the working directory, the data directory and the data directory's
    parent; the first that exists wins (else the first candidate, so the error message names it).
    """
    text = str(raw or "").strip().strip("\"'").strip()
    if not text:
        raise WorkbookError("workbook.path is not configured")
    path = Path(os.path.expandvars(text)).expanduser()
    if path.is_absolute():
        return path
    candidates = [Path.cwd() / path]
    if paths is not None:
        candidates += [paths.root / path, paths.root.parent / path]
    return next((c for c in candidates if c.exists()), candidates[0])


def read_workbook(
    path: Path,
    config: AppConfig,
    *,
    today: date,
    log: logging.Logger | None = None,
    streaming: bool | None = None,
) -> WorkbookParseResult:
    """Parse the workbook at ``path`` (see module docstring). Raises ``WorkbookError`` for source-wide problems.

    ``streaming=True`` forces openpyxl's read-only reader (no hyperlink objects / merged ranges); by default
    it is used only for very large files.
    """
    book = _open_book(path, streaming=streaming)
    try:
        selection = _select_sheet(book, config.workbook.sheet, config.workbook.column_map)
        return parse_sheet(selection.grid, selection.header, config, today=today, log=log)
    finally:
        book.close()


def _today_for(ctx: SourceContext) -> date:
    now = ctx.clock.now()
    try:
        return local_day(now, ctx.config.timezone)
    except (
        Exception
    ):  # unknown time zone name: use the UTC date rather than failing the whole source
        return now.date()


class WorkbookProvider:
    """``OpportunityProvider`` for the spreadsheet. ``fetch_report`` also says why rows were dropped."""

    name = "workbook"

    def enabled(self, config: AppConfig) -> bool:
        return bool(config.platforms.workbook and (config.workbook.path or "").strip())

    def fetch_report(self, ctx: SourceContext) -> WorkbookParseResult:
        path = resolve_workbook_path(ctx.config.workbook.path, ctx.paths)
        return read_workbook(path, ctx.config, today=_today_for(ctx), log=ctx.log)

    def fetch(self, ctx: SourceContext) -> list[Opportunity]:
        return self.fetch_report(ctx).opportunities


PROVIDER = WorkbookProvider()


# ------------------------------------------------------------------------------------------------ inspection


class SheetReport(BaseModel):
    name: str
    state: str = "visible"
    rows: int = 0
    columns: int = 0
    name_score: float = 0.0  # similarity to "verified opportunities"
    header_row: int | None = None
    recognised: list[str] = Field(default_factory=list)


class ColumnReport(BaseModel):
    column: str  # letter, "A".."Z", "AA"
    header: str
    maps_to: str | None = None  # canonical field; None -> stored in ``Opportunity.extra``


class WorkbookReport(BaseModel):
    """What `autoapply inspect-workbook` shows: sheets, the detected header, the column mapping, sample rows."""

    path: str
    sheets: list[str] = Field(default_factory=list)
    sheet_details: list[SheetReport] = Field(default_factory=list)
    sheet: str | None = None  # the sheet that would be ingested
    header_row: int | None = None  # 1-based
    headers: list[str] = Field(default_factory=list)
    mapping: dict[str, str] = Field(default_factory=dict)  # canonical field -> header text
    columns: list[ColumnReport] = Field(default_factory=list)
    sample_rows: list[dict[str, str]] = Field(default_factory=list)
    data_rows: int = 0
    kept: int = 0
    rejected: dict[str, int] = Field(default_factory=dict)
    rejected_rows: dict[str, list[str]] = Field(default_factory=dict)
    warnings: list[str] = Field(default_factory=list)

    def render(self) -> str:
        """Human-readable multi-line summary for the CLI."""
        lines = [f"Workbook: {self.path}", "Sheets: " + (", ".join(self.sheets) or "(none)")]
        if self.sheet is None:
            lines.append("No usable sheet / header row found.")
        else:
            lines.append(f"Using sheet {self.sheet!r}, header on row {self.header_row}")
            lines.append("Columns:")
            lines.extend(
                f"  {c.column:>3}  {c.header or '(blank)':<32} {c.maps_to or '-> extra'}"
                for c in self.columns
            )
            lines.append(f"Rows: {self.data_rows} with data, {self.kept} would be kept")
            lines.extend(
                f"  rejected {reason}: {count}" for reason, count in sorted(self.rejected.items())
            )
            if self.sample_rows:
                lines.append("Sample rows:")
                lines.extend(
                    "  " + "; ".join(f"{k}={v}" for k, v in row.items()) for row in self.sample_rows
                )
        lines.extend(f"Warning: {w}" for w in self.warnings)
        return "\n".join(lines)


def _sheet_reports(book: _Book, column_map: Mapping[str, str]) -> list[SheetReport]:
    reports: list[SheetReport] = []
    for name, state in book.sheets:
        grid = book.grid(name)
        header = detect_header(grid.rows, column_map)
        reports.append(
            SheetReport(
                name=name,
                state=state,
                rows=len(grid.rows),
                columns=len(grid.rows[0]) if grid.rows else 0,
                name_score=round(score_sheet_name(name), 1),
                header_row=header.row if header else None,
                recognised=sorted(header.fields) if header else [],
            )
        )
    return reports


def _sample_rows(grid: SheetGrid, header: HeaderInfo, size: int) -> list[dict[str, str]]:
    keys = [
        header.extras.get(i) or text or get_column_letter(i + 1)
        for i, text in enumerate(header.headers)
    ]
    samples: list[dict[str, str]] = []
    for row in grid.rows[header.row :]:
        if len(samples) >= size:
            break
        values = {
            keys[i]: clean_text(c.value, limit=120) or clean_text(c.link, limit=120)
            for i, c in enumerate(row)
            if i < len(keys)
        }
        if (
            len(filled := {k: v for k, v in values.items() if v}) >= 2
        ):  # single cells are separators
            samples.append(filled)
    return samples


def inspect_workbook(
    path: Path | str,
    config: AppConfig | None = None,
    *,
    today: date | None = None,
    sample_size: int = 5,
) -> WorkbookReport:
    """Diagnose a workbook without ingesting it: sheets, the detected header, column mapping, sample rows.

    ``config`` (default: defaults) supplies ``workbook.sheet`` / ``column_map`` and the search filters behind
    the kept / rejected counts; ``today`` defaults to the system date. Raises ``WorkbookError`` only when the
    file itself cannot be read; a missing header is reported (``sheet is None``) with the sheet list intact.
    """
    cfg = config or AppConfig()
    report = WorkbookReport(path=str(path))
    book = _open_book(Path(path))
    try:
        report.sheet_details = _sheet_reports(book, cfg.workbook.column_map)
        report.sheets = [s.name for s in report.sheet_details]
        try:
            selection = _select_sheet(book, cfg.workbook.sheet, cfg.workbook.column_map)
        except WorkbookError as exc:
            report.warnings.append(str(exc))
            return report
        header = selection.header
        by_col = {col: name for name, cols in header.fields.items() for col in cols}
        report.sheet = selection.grid.name
        report.header_row = header.row
        report.headers = list(header.headers)
        report.mapping = header.primary_headers()
        report.columns = [
            ColumnReport(column=get_column_letter(i + 1), header=text, maps_to=by_col.get(i))
            for i, text in enumerate(header.headers)
            if text or i in by_col
        ]
        report.sample_rows = _sample_rows(selection.grid, header, sample_size)
        parsed = parse_sheet(selection.grid, header, cfg, today=today or date.today())
        report.data_rows = parsed.data_rows
        report.kept = len(parsed.opportunities)
        report.rejected = parsed.rejection_counts()
        report.rejected_rows = parsed.rejection_summary()
        report.warnings.extend(parsed.warnings)
    finally:
        book.close()
    return report
