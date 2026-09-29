"""The sample workbook fixture (``autoapply.testing.fixtures``) and what the workbook reader makes of it."""

from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

from autoapply.config import AppConfig
from autoapply.models import ATS, Opportunity
from autoapply.normalize import canonical_url, fingerprint
from autoapply.sources.workbook import detect_ats, inspect_workbook, read_workbook
from autoapply.testing.fixtures import (
    SAMPLE_HEADERS,
    SAMPLE_SHEET,
    SAMPLE_TODAY,
    SITES,
    SampleRow,
    SampleUrls,
    SampleWorkbookInfo,
    build_sample_workbook,
)

FAMILIES = {
    "product_management",
    "technical_program_management",
    "technology_consulting",
    "strategy",
    "business_operations",
    "business_analysis",
    "analytics",
}


@pytest.fixture
def info(tmp_path: Path) -> SampleWorkbookInfo:
    return build_sample_workbook(tmp_path / "sample.xlsx")


# --------------------------------------------------------------------------------------------- the file itself


def test_workbook_layout(info: SampleWorkbookInfo) -> None:
    wb = load_workbook(info.path)
    assert wb.sheetnames == ["Summary", SAMPLE_SHEET, "Rejected"]
    ws = wb[SAMPLE_SHEET]
    assert ws["A1"].value.startswith("UT Austin - Verified Internship Opportunities")
    assert "A1:L1" in {
        str(r) for r in ws.merged_cells.ranges
    }  # a merged title row above the header
    assert [c.value for c in ws[4]] == list(SAMPLE_HEADERS)
    assert info.header_row == 4 and info.sheet == SAMPLE_SHEET and info.headers == SAMPLE_HEADERS
    summary = wb["Summary"]
    assert summary["A3"].value == "Metric"
    rejected = wb["Rejected"]
    assert [c.value for c in rejected[1]] == ["Company", "Role", "Link", "Reason"]
    assert rejected.max_row == 4
    wb.close()


def test_rows_are_written_where_the_info_says(info: SampleWorkbookInfo) -> None:
    wb = load_workbook(info.path)
    ws = wb[SAMPLE_SHEET]
    for row in info.rows:
        assert ws.cell(row=row.row_number, column=2).value == row.title, row.key
        company = ws.cell(row=row.row_number, column=1).value
        assert (company or "") == row.company, row.key
    wb.close()


def test_the_sample_covers_every_kind_of_link_and_date(info: SampleWorkbookInfo) -> None:
    wb = load_workbook(info.path)
    ws = wb[SAMPLE_SHEET]
    link_cells = [ws.cell(row=r.row_number, column=3) for r in info.rows]
    assert any(
        c.hyperlink is not None and c.value == "Apply" for c in link_cells
    )  # hyperlink over display text
    assert any(
        isinstance(c.value, str) and c.value.startswith("=HYPERLINK(") for c in link_cells
    )  # formula
    assert any(
        isinstance(c.value, str) and c.value.startswith("https://") for c in link_cells
    )  # plain text
    assert any(
        c.hyperlink is not None and str(c.value).startswith("https://short") for c in link_cells
    )
    assert any(c.value is None for c in link_cells)  # a row with no link at all
    verified = [ws.cell(row=r.row_number, column=8).value for r in info.rows]
    assert any(isinstance(v, int) for v in verified)  # Excel serial number
    assert any(isinstance(v, str) and "/" in v for v in verified)  # 9/20/26
    assert any(isinstance(v, str) and v[:4].isdigit() and "-" in v for v in verified)  # 2026-09-09
    assert any(isinstance(v, str) and v[:3].isalpha() for v in verified)  # Sep 18, 2026
    assert any(
        isinstance(v, str) and v[0].isdigit() and v.count(" ") == 2 for v in verified
    )  # 17 Sep 2026
    assert any(hasattr(v, "year") for v in verified)  # a real Excel date
    wb.close()


def test_ragged_blank_and_separator_rows_exist(info: SampleWorkbookInfo) -> None:
    ws = load_workbook(info.path)[SAMPLE_SHEET]
    ragged = info.row("Business Systems Analyst Intern")
    assert [c.value for c in ws[ragged.row_number]][4:] == [None] * 8  # nothing beyond the location
    lines = [[c.value for c in row] for row in ws.iter_rows(min_row=5)]
    assert any(all(v is None for v in line) for line in lines)  # blank rows
    assert any(
        line[0] and all(v is None for v in line[1:]) for line in lines
    )  # separator / footer rows


def test_the_sample_is_deterministic(tmp_path: Path) -> None:
    a = build_sample_workbook(tmp_path / "a.xlsx")
    b = build_sample_workbook(tmp_path / "b.xlsx")
    assert [(r.key, r.url, r.id, r.row_number) for r in a.rows] == [
        (r.key, r.url, r.id, r.row_number) for r in b.rows
    ]
    assert a.expected_ids == b.expected_ids
    assert a.expected_rejections == b.expected_rejections


def test_builder_creates_missing_folders_and_accepts_str(tmp_path: Path) -> None:
    target = tmp_path / "deep" / "er" / "Verified — sample (1).xlsx"
    built = build_sample_workbook(str(target))
    assert target.is_file() and built.path == target


# --------------------------------------------------------------------------------------------- expectations


def test_expectations_are_internally_consistent(info: SampleWorkbookInfo) -> None:
    assert len({r.key for r in info.rows}) == len(info.rows)
    assert len({r.row_number for r in info.rows}) == len(info.rows)
    assert info.expected_ids and len(set(info.expected_ids)) == len(info.expected_ids)
    dropped = {t for titles in info.expected_rejections.values() for t in titles}
    assert dropped.isdisjoint({r.title for r in info.rows if r.outcome == "ingested"})
    assert set(info.expected_rejections) == {
        "closed",
        "deadline_passed",
        "wrong_term",
        "stale",
        "not_internship",
        "no_url",
        "missing_fields",
    }
    assert info.expected_rejections["closed"] == [
        "Strategic Finance Intern",
        "Process Analyst Intern",
        "Data Analyst Intern",
        "Product Operations Intern",
    ]
    assert info.expected_rejections["wrong_term"] == [
        "Corporate Development Intern",
        "Supply Chain Analyst Intern",
    ]
    assert set(info.expected_duplicates) <= set(info.expected_ids)
    assert sum(len(v) for v in info.expected_duplicates.values()) == 4
    assert [r.row_number for r in info.rows] == sorted(r.row_number for r in info.rows)
    assert len(info.workbook_kept) == len(info.expected_ids) + 4


def test_sample_size_and_coverage(info: SampleWorkbookInfo) -> None:
    assert 30 <= len(info.rows) <= 40
    assert len(info.expected_ids) == 15
    assert {r.family for r in info.eligible} == FAMILIES  # every README role family
    for family in FAMILIES:
        assert len([r for r in info.eligible if r.family == family]) >= 2, family
    assert {r.site for r in info.eligible} == set(SITES)  # every logical mock site
    assert [r.title for r in info.rows if r.outcome == "ingested" and r.family is None] == [
        "Marketing Intern"
    ]
    assert info.row("Senior Product Manager").outcome == "not_internship"
    outcomes = {r.outcome for r in info.rows}
    assert {
        "ingested",
        "duplicate",
        "closed",
        "wrong_term",
        "stale",
        "deadline_passed",
        "not_internship",
    } <= outcomes


def test_row_lookup_helpers(info: SampleWorkbookInfo) -> None:
    assert info.row("Marketing Intern").company == "Elmwood Retail"
    assert info.row("Business Analyst Intern", "Kestrel Aerospace").site == "workday"
    with pytest.raises(KeyError):
        info.row("No such role")
    survivor = info.row(
        "Product Management Intern"
    )  # listed twice: the tracking-parameter copy is a duplicate
    assert survivor.outcome == "ingested" and survivor.key == "alder_pm"
    assert info.by_key("alder_pm_dup").duplicate_of == survivor.id
    with pytest.raises(KeyError):
        info.row("Product Management Intern", "Nobody Inc")
    with pytest.raises(KeyError):
        info.by_key("nope")
    assert {r.title for r in info.by_site("workday")} == {
        "Product Management Intern",
        "Technology Consulting Summer Analyst",
        "Business Analyst Intern",
    }
    with pytest.raises(KeyError):
        SampleUrls().for_site("nope")


# --------------------------------------------------------------------------------------------- the reader on the sample


def test_reader_matches_the_expectations(info: SampleWorkbookInfo) -> None:
    result = read_workbook(info.path, AppConfig(), today=info.today)
    assert result.sheet == SAMPLE_SHEET and result.header_row == 4
    assert result.rejection_summary() == info.expected_rejections
    assert [o.title for o in result.opportunities] == [r.title for r in info.workbook_kept]
    assert [o.id for o in result.opportunities] == [r.id for r in info.workbook_kept]
    assert result.data_rows == len(info.rows)
    assert result.junk_rows == 6  # three section rows, two blank rows, the footer
    assert result.warnings == []


def test_every_surviving_row_is_read_faithfully(info: SampleWorkbookInfo) -> None:
    result = read_workbook(info.path, AppConfig(), today=info.today)
    by_key = {r.row_number: r for r in info.rows}
    for opp in result.opportunities:
        row = by_key[opp.extra["sheet_row"]]
        assert opp.company == row.company and opp.title == row.title, row.key
        assert opp.url == row.url, row.key
        assert opp.apply_url is None
        assert opp.term == "Summer 2027"
        assert opp.last_verified == row.last_verified, row.key
        assert opp.posted_date == row.posted, row.key
        assert opp.deadline == row.deadline, row.key
        assert opp.ats == row.ats or row.outcome == "duplicate", row.key
        assert opp.is_open
        for flag in ("term_assumed", "date_unknown"):
            assert bool(opp.extra.get(flag)) is (flag in row.flags), (row.key, flag)
        if row.location and row.key != "lake_bsa":
            assert opp.location == row.location, row.key


def test_raw_columns_land_in_extra(info: SampleWorkbookInfo) -> None:
    result = read_workbook(info.path, AppConfig(), today=info.today)
    alder = next(
        o
        for o in result.opportunities
        if o.title == "Product Management Intern" and "?" not in o.url
    )
    assert alder.extra["Pay (hourly)"] == "$34"
    assert alder.extra["notes"] == "Referral welcome"
    assert alder.extra["status_raw"] == "Open"
    assert alder.extra["section"] == "Product & Program Management"
    long_notes = next(o for o in result.opportunities if o.title == "Data Analytics Intern")
    assert long_notes.description is not None and len(long_notes.description) > 10_000
    assert len(long_notes.extra["notes"]) <= 1000 and long_notes.extra["notes"].endswith(
        "\u2026"
    )  # capped
    assert (
        long_notes.extra["term_raw"] == "Summer 2027"
    )  # the "applications open Fall 2026" note did not matter


def test_status_vocabulary_of_the_sample_is_recorded(info: SampleWorkbookInfo) -> None:
    result = read_workbook(info.path, AppConfig(), today=info.today)
    statuses = {o.extra.get("status_raw") for o in result.opportunities}
    assert {"Open", "Active", "Verified", "Yes", "Live"} <= statuses


def test_the_default_urls_are_placeholders_and_ats_comes_from_the_ats_column(
    info: SampleWorkbookInfo,
) -> None:
    for row in info.rows:
        assert row.url == "" or "example.test" in row.url or row.site == "aggregator"
    result = read_workbook(info.path, AppConfig(), today=info.today)
    by_title: dict[tuple[str, str], Opportunity] = {}
    for opp in result.opportunities:
        by_title.setdefault(
            (opp.company, opp.title), opp
        )  # the first listing of a role is the survivor
    assert by_title[("Alder Systems", "Product Management Intern")].ats == ATS.WORKDAY
    assert by_title[("Fernhill Software", "Technology Consultant Intern")].ats == ATS.CUSTOM
    assert (
        by_title[("Kestrel Aerospace", "Business Analyst Internship (Summer 2027)")].ats
        == ATS.UNKNOWN
    )


def test_alternative_view_through_inspect_workbook(info: SampleWorkbookInfo) -> None:
    report = inspect_workbook(info.path, AppConfig(), today=info.today)
    assert report.sheets == ["Summary", SAMPLE_SHEET, "Rejected"]
    assert report.sheet == SAMPLE_SHEET and report.header_row == 4
    assert set(report.mapping) == {
        "company",
        "title",
        "apply_url",
        "location",
        "term",
        "status",
        "posted",
        "verified",
        "deadline",
        "ats",
        "notes",
    }
    assert report.kept == len(info.workbook_kept)
    assert report.rejected_rows == info.expected_rejections
    assert report.columns[-1].header == "Pay (hourly)" and report.columns[-1].maps_to is None
    assert len(report.sample_rows) == 5
    assert report.warnings == []


# --------------------------------------------------------------------------------------------- customising the sample


def test_custom_urls_and_today(tmp_path: Path) -> None:
    urls = SampleUrls(
        workday="http://acme.wd5.myworkdayjobs.com.localhost:5001",
        greenhouse="http://boards.greenhouse.io.localhost:5002/",
        lever="http://jobs.lever.co.localhost:5003",
        ashby="http://jobs.ashbyhq.com.localhost:5004",
        portal="http://careers.example-employer.com.localhost:5005",
        captcha="http://blockers.example.localhost:5006/captcha/job-1",
        closed="http://blockers.example.localhost:5006/closed/job-2",
        sso="http://blockers.example.localhost:5006/sso/job-3",
    )
    later = SAMPLE_TODAY + timedelta(days=400)
    info = build_sample_workbook(tmp_path / "custom.xlsx", urls, today=later)
    assert info.today == later and info.urls == urls
    result = read_workbook(info.path, AppConfig(), today=later)
    assert result.rejection_summary() == info.expected_rejections
    assert [o.id for o in result.opportunities] == [r.id for r in info.workbook_kept]
    for row in info.rows:
        if row.site in ("workday", "greenhouse", "lever", "ashby"):
            assert row.url.startswith(urls.for_site(row.site).rstrip("/") + row.path), row.key
            assert row.path.startswith("/")
        elif row.site in ("captcha", "closed", "sso"):
            assert (
                row.url == urls.for_site(row.site) and row.path == ""
            )  # base with a path is used verbatim
    by_key = {r.key: r for r in info.rows}
    assert (
        by_key["fern_tc"].url
        == "http://careers.example-employer.com.localhost:5005/careers/jobs/technology-consultant-intern"
    )
    # real-looking mock hosts: the ATS is now recognisable from the URL alone
    for opp in result.opportunities:
        row = next(r for r in info.rows if r.row_number == opp.extra["sheet_row"])
        if row.site in ("workday", "greenhouse", "lever", "ashby"):
            assert detect_ats(opp.url) == row.ats == opp.ats, row.key
    # dates are relative to the chosen day
    kept = {o.extra["sheet_row"]: o for o in result.opportunities}
    alder = kept[info.row("Product Management Intern").row_number]
    assert alder.last_verified == later - timedelta(days=5)


def test_urls_use_real_ats_layouts(info: SampleWorkbookInfo) -> None:
    urls = SampleUrls(
        workday="https://acme.wd5.myworkdayjobs.com",
        greenhouse="https://boards.greenhouse.io",
        lever="https://jobs.lever.co",
        ashby="https://jobs.ashbyhq.com",
    )
    real = build_sample_workbook(info.path.with_name("real.xlsx"), urls)
    workday = real.row("Product Management Intern")
    assert workday.path.startswith("/en-US/External/job/Austin-TX/Product-Management-Intern_R-")
    assert workday.job_id.startswith("R-")
    greenhouse = real.row("Associate Product Manager Intern")
    assert (
        greenhouse.path == f"/birchwood-health/jobs/{greenhouse.job_id}"
        and greenhouse.job_id.isdigit()
    )
    for title in ("Technical Program Manager Intern", "Program Management Intern"):  # Lever, Ashby
        row = real.row(title)
        company_slug, job = row.path.strip("/").split("/")
        assert len(job) == 36 and job.count("-") == 4  # a uuid
        assert company_slug in row.url and row.job_id == job
    for row in real.rows:
        if row.site in ("workday", "greenhouse", "lever", "ashby"):
            assert detect_ats(row.url) == row.ats, row.key


def test_duplicate_rows_collapse_the_way_the_expectations_say(info: SampleWorkbookInfo) -> None:
    by_key = {r.key: r for r in info.rows}
    # tracking-parameter variants share the survivor's id
    for dup, survivor in (("granite_strat_dup", "granite_strat"), ("alder_pm_dup", "alder_pm")):
        assert by_key[dup].url != by_key[survivor].url
        assert canonical_url(by_key[dup].url) == canonical_url(by_key[survivor].url)
        assert by_key[dup].id == by_key[survivor].id == by_key[dup].duplicate_of
    # the same role under a different URL only shares the fingerprint, and the direct ATS row must win
    for dup, survivor in (("kestrel_ba_linkedin", "kestrel_ba"), ("iron_biz_indeed", "iron_biz")):
        a, b = by_key[dup], by_key[survivor]
        assert a.id != b.id and a.duplicate_of == b.id
        assert fingerprint(a.company, a.title, a.location) == fingerprint(
            b.company, b.title, b.location
        )
        assert b.site in ("workday", "ashby") and a.site == "aggregator"
    assert (
        by_key["iron_biz_indeed"].row_number < by_key["iron_biz"].row_number
    )  # aggregator listed first


def test_stale_rows_are_older_than_the_window(info: SampleWorkbookInfo) -> None:
    cutoff = info.today - timedelta(days=45)
    for row in info.rows:
        newest = max((d for d in (row.last_verified, row.posted) if d), default=None)
        if row.outcome == "stale":
            assert newest is not None and newest < cutoff, row.key
        elif row.outcome == "ingested" and newest is not None:
            assert newest >= cutoff, row.key
    deadline_row = info.row("Technology Advisory Intern")
    assert deadline_row.deadline is not None and deadline_row.deadline < info.today


def test_sample_row_is_a_frozen_value_object(info: SampleWorkbookInfo) -> None:
    row: SampleRow = info.rows[0]
    with pytest.raises(AttributeError):
        row.title = "changed"  # type: ignore[misc]
    assert isinstance(row.last_verified, date | type(None))
